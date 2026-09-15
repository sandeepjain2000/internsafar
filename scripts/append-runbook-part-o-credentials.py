#!/usr/bin/env python3
"""Append Part O — separately provided credentials / do not reuse old .env."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
MARKER = "Part O — New production credentials (provided separately — not in the zip)"


def main():
    doc = Document(DOC_PATH)
    existing = "\n".join(p.text for p in doc.paragraphs)
    if MARKER in existing:
        print("Part O already present — skip")
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)

    doc.add_paragraph(
        "The handoff zip does NOT contain .env, Google client secret JSON, or any real secrets. "
        "The person handing off this pack will give you the NEW production values separately "
        "(message/email/secure channel — outside the zip)."
    )

    doc.add_heading("Do not reuse old EC2 .env values", level=2)
    for line in [
        "Old or leftover NEXTAUTH_URL (especially http://localhost:3000) caused sign-out and Google OAuth to redirect to localhost.",
        "Missing or wrong NEXTAUTH_SECRET caused login/CAPTCHA failures in production.",
        "Stale GOOGLE_CLIENT_ID / GOOGLE_CLIENT_SECRET (or wrong redirect URIs) break Google registration verification.",
        "Do NOT copy a developer laptop .env.local onto EC2.",
        "Do NOT keep an old EC2 .env unchanged when you were given new credentials for this deploy.",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_heading("What to put on EC2", level=2)
    doc.add_paragraph(
        "On the EC2 host, edit ~/internship-portal/.env and set (or replace) these using the "
        "NEW values provided separately to you:"
    )
    for line in [
        "NEXTAUTH_URL=https://internsafar.com",
        "NEXTAUTH_SECRET=<new secret from the handoff person — or generate with: openssl rand -base64 32>",
        "GOOGLE_CLIENT_ID=<from the separately provided Google OAuth / client credentials>",
        "GOOGLE_CLIENT_SECRET=<from the separately provided Google OAuth / client credentials>",
        "DATABASE_URL=<keep working RDS URL unless you were given a new one>",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_paragraph(
        "If you were given a Google Cloud client secret JSON file separately: open it only to copy "
        "client_id and client_secret into .env. Do not commit that file, do not leave it in the "
        "handoff folder, and do not upload it into GitHub."
    )

    doc.add_heading("Apply after editing .env", level=2)
    for line in [
        "cd ~/internship-portal",
        "npm run build",
        "pm2 restart internsafar --update-env",
        "Verify: sign-out must NOT go to localhost; Google redirect URI in Cloud Console must be https://internsafar.com/api/auth/callback/google",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_paragraph(
        "Same warning is in the handoff README.txt and AWS-DEPLOY.md."
    )

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
