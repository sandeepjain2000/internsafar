# -*- coding: utf-8 -*-
"""
Add TC-IS-06-011 (candidate name fields reject digits / symbols) to
InternSafar-Test-Cases.xlsx — from external QA report BUG-PROFILE-001.

Rule (src/lib/ipPersonName.js): letters in any script, spaces, hyphens, apostrophes,
full stops; must start with a letter. Checked inline on /candidate/profile and by
PUT /api/ip/candidate/profile (400). Automated by Playwright regression IS-079.

  python scripts/sync-internsafar-xlsx-2026-10-07-name-validation.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-07-name-validation.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "06 Candidate Profile"
TEMPLATE_ID = "TC-IS-06-010"

NEW_CASE = {
    "ID": "TC-IS-06-011",
    "Module / Section": "06 Candidate Profile",
    "Type": "Negative",
    "Issue Summary": "First / Middle / Last Name reject digits and symbols",
    "Description": (
        "1. Sign in as the test candidate and open /candidate/profile → Basics.\n"
        "2. Type 'Priya123' in First Name (also try '123', 'Priya@', Middle 'K2', Last 'Sharma_').\n"
        "3. Click Save.\n"
        "4. Replace with a valid name such as Anne-Marie O'Neil, José or R. K."
    ),
    "Suggestion / Expected Behaviour": (
        "An inline message appears under the field: '<Field> can only contain letters, spaces, hyphens (-), "
        "apostrophes (') and full stops (.).' Save shows 'Not saved — …' and nothing is stored. "
        "PUT /api/ip/candidate/profile with an invalid name returns 400 with the same message. "
        "Letters in any language, spaces, hyphens, apostrophes and full stops are accepted; names must start with a letter."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P1",
    "Test Status": "Not Run",
    "Comments / Notes": "From external QA report BUG-PROFILE-001 (2026-10-07). Playwright IS-079.",
    "Phase": "2026-10-07 name validation",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Profile name validation",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        print(f"{NEW_CASE['ID']} already present (row {rows[NEW_CASE['ID']]}) — nothing to add.")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
