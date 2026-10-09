# -*- coding: utf-8 -*-
"""
Add TC-IS-06-014 (candidate profile: labels tied to inputs, preferred hours / commitment note rules,
save message names the tab of the problem field) to InternSafar-Test-Cases.xlsx.

Rules: src/lib/ipProfileSetup.js — preferred hours only need to be valid times (night shifts such as
22:00–02:00 are allowed); commitment note 200 characters max; only changed values are checked. Same
rules on /candidate/profile (inline) and PUT /api/ip/candidate/profile (400). Automated by Playwright
IS-082. Re-running with --apply updates the row if it already exists.

  python scripts/sync-internsafar-xlsx-2026-10-08-profile-a11y-hours.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-08-profile-a11y-hours.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "06 Candidate Profile"
TEMPLATE_ID = "TC-IS-06-013"

NEW_CASE = {
    "ID": "TC-IS-06-014",
    "Module / Section": "06 Candidate Profile",
    "Type": "Negative",
    "Issue Summary": "Profile labels are tied to inputs; night-shift hours accepted; commitment note limited; save message names the tab",
    "Description": (
        "1. Sign in as the test candidate and open /candidate/profile.\n"
        "2. Click the 'First Name' label — the First Name box gets focus. Type 'Priya2' (screen reader reads the error with the field).\n"
        "3. 4. Work Readiness → Preferred working hours: from 22:00, to 02:00 (night shift). Save.\n"
        "4. Ongoing commitment? = Other → Commitment note: try to type more than 200 characters.\n"
        "5. Open the Privacy & Photo tab → Telegram '@ab'. Go back to 1. Basics & Contact and click Save.\n"
        "6. Click 'Go to Privacy & Photo' in the save message."
    ),
    "Suggestion / Expected Behaviour": (
        "Step 2: label focuses the input; the name error is announced with it. "
        "Step 3: no error — night-shift hours (\"to\" earlier than \"from\") save as entered. "
        "Step 4: the box stops at 200 characters (API also returns 400 above 200). "
        "Step 5: 'Not saved — Telegram handle must be … (on the Privacy & Photo tab)' with a 'Go to Privacy & Photo' link. "
        "Step 6: the Privacy tab opens and the Telegram box is focused. "
        "PUT /api/ip/candidate/profile returns 400 only when an hours value is not a time (e.g. '9am') or the note is over 200 characters."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P2",
    "Test Status": "Not Run",
    "Comments / Notes": "From external QA report recommendations (2026-10-08). Playwright IS-082.",
    "Phase": "2026-10-08 profile labels and hours",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Profile accessibility / setup validation",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        text_keys = ("Issue Summary", "Description", "Suggestion / Expected Behaviour", "Feature")
        for key in text_keys:
            if key in cols:
                ws.cell(row, cols[key]).value = NEW_CASE[key]
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
