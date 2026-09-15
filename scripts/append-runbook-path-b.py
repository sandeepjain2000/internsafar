#!/usr/bin/env python3
"""Append Part M — Path B pastable app-update section to CRISP (1)(1) runbook."""
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"

MARKER = "Part M — Path B (app update on live EC2) — pastable"


def add_mono(doc, text):
    """One command/line as monospace-friendly plain paragraph (pasteable)."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    return p


def main():
    doc = Document(DOC_PATH)
    existing = "\n".join(p.text for p in doc.paragraphs)
    if MARKER in existing:
        print("Part M already present — skip")
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)

    doc.add_paragraph(
        "Use this section when EC2 + RDS + Nginx are already live and you only need to "
        "deploy a new internship-portal-aws-handoff.zip. Do NOT run deploy:fresh-aws-db "
        "(that resets/seeds the database — Path C only)."
    )
    doc.add_paragraph(
        "Prerequisites: handoff zip already SCP’d and unzipped to ~/internship-portal-aws-handoff "
        "(contains internship-portal-aws-deploy-*.tar.gz). A live app exists at ~/internship-portal "
        "with a working .env."
    )

    doc.add_heading("Path B — paste block 1: extract + permissions", level=2)
    doc.add_paragraph("Paste this entire block into the EC2 SSH session:")
    for line in [
        "cd ~/internship-portal-aws-handoff",
        "mkdir -p ~/internship-portal-new",
        "tar -xzf internship-portal-aws-deploy-*.tar.gz -C ~/internship-portal-new",
        "chmod -R u+rwX ~/internship-portal-new",
        "chown -R ubuntu:ubuntu ~/internship-portal-new",
        "tar -tzf internship-portal-aws-deploy-*.tar.gz | head",
        "ls ~/internship-portal-new/internship-portal/package.json",
    ]:
        add_mono(doc, line)

    doc.add_heading("Path B — paste block 2: app-swap (preserve .env)", level=2)
    doc.add_paragraph(
        "Paste this entire block. It stops on the first failure (set -e). "
        "Source path MUST be ~/internship-portal-new/internship-portal (not the handoff root)."
    )
    for line in [
        "set -e",
        "cp ~/internship-portal/.env ~/internship-portal.env.backup",
        "mv ~/internship-portal ~/internship-portal-old-$(date +%Y%m%d)",
        "mv ~/internship-portal-new/internship-portal ~/internship-portal",
        "cp ~/internship-portal.env.backup ~/internship-portal/.env",
        "chmod 600 ~/internship-portal/.env",
        "cd ~/internship-portal",
        "grep -E '^(DATABASE_|NEXTAUTH_|GOOGLE_)' .env",
    ]:
        add_mono(doc, line)

    doc.add_heading("Path B — paste block 3: install, build, PM2", level=2)
    doc.add_paragraph("Paste this entire block after the app-swap succeeds:")
    for line in [
        "cd ~/internship-portal",
        "npm install --legacy-peer-deps",
        "npm run build",
        "pm2 restart internsafar --update-env",
        "pm2 status",
        "curl -sI https://internsafar.com | head",
    ]:
        add_mono(doc, line)

    doc.add_heading("Path B — verify (do not skip)", level=2)
    for line in [
        "Confirm NEXTAUTH_URL=https://internsafar.com in ~/internship-portal/.env (not localhost).",
        "Sign out as candidate/employer → must land on https://internsafar.com/ (not localhost:3000).",
        "If .env was changed: npm run build THEN pm2 restart internsafar --update-env.",
        "Do NOT run npm run deploy:fresh-aws-db on Path B unless you intend to reset RDS data.",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_paragraph(
        "Same Path B pastable blocks are also in the handoff zip file AWS-DEPLOY.md "
        "(section “Path B — full pastable”)."
    )

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH} ({len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    main()
