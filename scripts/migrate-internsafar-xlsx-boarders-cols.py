# -*- coding: utf-8 -*-
"""Migrate InternSafar-Test-Cases.xlsx to Boarders column names + IP extras.

Preserves row data. Run once (idempotent if already migrated).

  python scripts/migrate-internsafar-xlsx-boarders-cols.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib.internsafar_xlsx_cols import COLS, OLD_TO_NEW

XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SKIP = frozenset({"Index", "Coverage Matrix", "Coverage", "Notes", "How to use", "Meta"})

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
BODY_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")


def find_header(ws):
    for r in range(1, 12):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 30) + 1)]
        if not any(vals):
            continue
        if "ID" in vals and "Issue Summary" in vals and "Test Status" in vals:
            return r, {str(v): i + 1 for i, v in enumerate(vals) if v}, True
        if "TC ID" in vals and "Status" in vals:
            return r, {str(v): i + 1 for i, v in enumerate(vals) if v}, False
    return None, {}, False


def migrate_sheet(ws) -> str:
    hr, cols, already = find_header(ws)
    if not hr:
        return "skip-no-header"
    if already and all(k in cols for k in COLS):
        return "ok-already"

    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, max(ws.max_column, 1) + 1)):
            continue
        rec = {name: ws.cell(r, col).value for name, col in cols.items()}
        out = {k: None for k in COLS}
        for old, new in OLD_TO_NEW.items():
            if old in rec and rec[old] is not None:
                out[new] = rec[old]
        for k in COLS:
            if k in rec and out.get(k) is None:
                out[k] = rec[k]
        if not out.get("Doc Status"):
            out["Doc Status"] = "Ready for QA"
        if not out.get("Reference") and out.get("Automation"):
            out["Reference"] = out.get("Automation")
        if not out.get("ID") and not out.get("Issue Summary"):
            continue
        rows.append(out)

    ws.delete_rows(hr, ws.max_row - hr + 1)

    for c, name in enumerate(COLS, 1):
        cell = ws.cell(hr, c, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, rec in enumerate(rows):
        r = hr + 1 + i
        for c, name in enumerate(COLS, 1):
            cell = ws.cell(r, c, rec.get(name))
            cell.font = BODY_FONT
            cell.alignment = WRAP

    for i, w in enumerate(
        [12, 28, 14, 42, 40, 40, 28, 14, 14, 12, 16, 28, 16, 20, 16, 10, 18, 28, 28, 22, 18, 14],
        1,
    ):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hr + 1}"
    return f"migrated:{len(rows)}"


def main():
    wb = load_workbook(XLSX)
    summary = {}
    for ws in wb.worksheets:
        if ws.title in SKIP:
            summary[ws.title] = "skipped"
            continue
        summary[ws.title] = migrate_sheet(ws)
    if "Index" in wb.sheetnames:
        idx = wb["Index"]
        note = (
            "Column schema: Boarders checklist names + Role(s)/Preconditions/"
            "Actual Result/Automation/Feature/Legacy ID."
        )
        found = any(
            idx.cell(r, 1).value and "Boarders checklist names" in str(idx.cell(r, 1).value)
            for r in range(1, min(idx.max_row, 50) + 1)
        )
        if not found:
            idx.cell(idx.max_row + 2, 1, note).font = BODY_FONT
    wb.save(XLSX)
    print(f"Saved {XLSX}")
    print(summary)


if __name__ == "__main__":
    main()
