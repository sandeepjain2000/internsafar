#!/usr/bin/env python3
"""Append Section B/C deployment lessons to CRISP (1)(1) runbook."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"


def add_para(doc, text, bullet=False, number=None):
    prefix = ""
    if bullet:
        prefix = "• "
    elif number is not None:
        prefix = f"{number}. "
    return doc.add_paragraph(prefix + text)


def main():
    doc = Document(DOC_PATH)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("Part L — Deployment lessons (September 2026)", level=1)
    add_para(doc, "This section supplements Parts A–K with lessons from the September 2026 AWS deployment session. "
              "It does not replace or rewrite earlier steps.")

    doc.add_heading("Three deployment paths", level=2)
    rows = [
        ("A — First-time AWS", "New EC2 + new RDS", "Parts A–K Steps 0–51"),
        ("B — App update on live EC2", "New code drop, RDS unchanged",
         "SCP handoff zip → unzip → chmod/chown → app-swap block → npm install → build → pm2 restart --update-env"),
        ("C — Fresh database on existing EC2", "RDS empty or intentionally reset",
         "Step 22 RDS test → IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db → verify demo accounts → build/PM2"),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Path"
    hdr[1].text = "When"
    hdr[2].text = "Key steps"
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val

    doc.add_heading("Step 18b — Handoff tar extract and permissions", level=2)
    add_para(doc, "After SCP of internship-portal-aws-handoff.zip to EC2 (handoff folder contains a .tar.gz app archive, not a .zip):")
    for line in [
        "cd ~/internship-portal-aws-handoff",
        "mkdir -p ~/internship-portal-new",
        "tar -xzf internship-portal-aws-deploy-*.tar.gz -C ~/internship-portal-new",
        "chmod -R u+rwX ~/internship-portal-new",
        "chown -R ubuntu:ubuntu ~/internship-portal-new",
        "tar -tzf internship-portal-aws-deploy-*.tar.gz | head",
        "# App-swap uses: ~/internship-portal-new/internship-portal",
    ]:
        add_para(doc, line, bullet=True)

    doc.add_heading("Step 23 (updated) — Database setup for fresh RDS", level=2)
    for n, line in enumerate([
        "npm install --legacy-peer-deps",
        "RDS connection test (Step 22)",
        "IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db  (001–034 → IP_Reset_Core_Sample.js → 035–039)",
        "Confirm === OK === banners and exit code 0 before continuing (=== BLOCKED === means allow env missing)",
    ], start=1):
        add_para(doc, line, number=n)

    add_para(doc, "After 001–034 on a fresh RDS, ip_users may show only SuperAdmin — this is expected until core seed runs.")

    add_para(doc, "Demo account verification (password Admin@123 for all):")
    for cred in [
        "Candidate: lawsonlclintern+1@gmail.com",
        "Employer: placementhubsupport@gmail.com",
        "SuperAdmin: support@placementhub.online",
    ]:
        add_para(doc, cred, bullet=True)

    doc.add_heading("Step 23d — Google OAuth production setup", level=2)
    for line in [
        "Set NEXTAUTH_URL=https://internsafar.com in ~/internship-portal/.env",
        "Google Cloud Console — Authorized JavaScript origin: https://internsafar.com",
        "Google Cloud Console — Redirect URI: https://internsafar.com/api/auth/callback/google",
        "npm run build (required after .env change)",
        "pm2 restart internsafar --update-env",
        "Symptom: browser redirects to localhost:3000/api/auth/callback/google → fix NEXTAUTH_URL and rebuild",
        "Note: Google verifies registration only; login remains email + password",
    ]:
        add_para(doc, line, bullet=True)

    doc.add_heading("Step 23e — Auth verification (sign-out + OAuth)", level=2)
    for line in [
        "Candidate/employer sign-out must land on https://internsafar.com/ (not localhost:3000)",
        "SuperAdmin sign-out must land on https://internsafar.com/superadmin/login",
        "If sign-out or Google OAuth redirects to localhost → check NEXTAUTH_URL, rebuild, PM2 --update-env",
        "Account page Active sessions should drop current session after sign-out",
    ]:
        add_para(doc, line, bullet=True)

    doc.add_heading("Step 26 (extend) — PM2 and Server Action logs", level=2)
    add_para(doc, "Always use: pm2 restart internsafar --update-env after .env changes.")
    add_para(doc, "PM2 log line 'The Server Reference ID did not match the expected format' usually means stale "
              "browser JS vs server build. If HTTP 200 and the site works in the browser, treat as log-only. "
              "If UI breaks: npm run build → pm2 restart --update-env → hard refresh (Ctrl+Shift+R).")

    doc.add_heading("Step 40 (extend) — NEXTAUTH_SECRET and captcha", level=2)
    for line in [
        "Missing NEXTAUTH_SECRET breaks sign-in and captcha in production",
        "Adding NEXTAUTH_SECRET to .env alone does not apply it — npm run build required before PM2 restart",
        "If .env was modified after .next was built, rebuild before PM2 restart (check timestamps)",
        "Interactive diagnosis: pm2 flush → try sign-in → pm2 logs for [next-auth][error][NO_SECRET]",
    ]:
        add_para(doc, line, bullet=True)

    doc.add_heading("Step 41 (extend) — RDS SSL vs auth rebuild", level=2)
    for line in [
        "SSL CA change (DATABASE_SSL_CA) is runtime-only — pm2 restart --update-env is enough; no rebuild",
        "Verify RDS CA bundle ~160 KB (162K): /etc/ssl/rds/global-bundle.pem",
        "After SSL fix, 'relation does not exist' means DB connected but schema wrong → Part I if needed",
    ]:
        add_para(doc, line, bullet=True)

    doc.add_heading("Step 28 (note) — Nginx server_name", level=2)
    add_para(doc, "server_name is an Nginx config directive inside /etc/nginx/sites-available/ — not a shell command.")

    doc.add_heading("Tool routing (migrations)", level=2)
    table2 = doc.add_table(rows=4, cols=2)
    table2.rows[0].cells[0].text = "Tool"
    table2.rows[0].cells[1].text = "When"
    tools = [
        ("IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db", "Path C — intentional fresh/empty RDS"),
        ("node recreate.mjs in ~/internsafar-aws", "Part I only — missing tables emergency (destructive)"),
        ("Path B — app code update only", "Do not run any DB migrate/seed; do not set IP_ALLOW_DB_MIGRATE"),
    ]
    for i, (a, b) in enumerate(tools, start=1):
        table2.rows[i].cells[0].text = a
        table2.rows[i].cells[1].text = b

    add_para(doc, "Migration manifest: 40 SQL files (001–039, includes two 008_* files).")

    doc.add_heading("Handoff runner (issues #7–8)", level=2)
    for line in [
        "Handoff runner needs pg from app node_modules: NODE_PATH=~/internship-portal/node_modules",
        "Copy .env: cp ~/internship-portal/.env ~/internship-portal-aws-handoff/.env",
        "Or run migrations from ~/internship-portal with IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db",
    ]:
        add_para(doc, line, bullet=True)

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH} ({len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    main()
