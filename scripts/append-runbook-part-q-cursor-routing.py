#!/usr/bin/env python3
"""Append Part Q — Cursor / agent DB script routing (current commands only)."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
MARKER = "Part Q — Agent DB routing (current)"


def main():
    doc = Document(DOC_PATH)
    if MARKER in "\n".join(p.text for p in doc.paragraphs):
        print("Part Q already present — skip")
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)
    doc.add_paragraph(
        "Agents (including Cursor) must use the table below — do not invent alternate migrate commands."
    )

    doc.add_heading("Use this table", level=2)
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Situation"
    table.rows[0].cells[1].text = "Command"
    rows = [
        ("Fresh / empty AWS RDS (Path C)", "npm run deploy:fresh-aws-db"),
        ("SQL only; demo users already exist", "npm run db:migrate:sql-only"),
        ("Path B — app code update only", "Do not run any DB migrate/seed"),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b

    doc.add_paragraph(
        "SQL-only file: scripts/db_migrate_sql_only_ip.mjs "
        "(handoff runner: runner/db_migrate_sql_only_ip.mjs)."
    )
    doc.add_paragraph(
        "Also documented in internship-portal/AGENTS.md (AWS / RDS database scripts), "
        "README.txt, AWS-DEPLOY.md, and DB-REFERENCE.md."
    )
    doc.add_paragraph(
        "Success requires === OK === banners and exit code 0. === FAIL === means stop."
    )

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
