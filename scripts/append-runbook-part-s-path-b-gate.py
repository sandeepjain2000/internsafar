#!/usr/bin/env python3
"""Append/update Part S — Path B no-DB gate (why / how / code)."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
MARKER = "Part S — Path B must not run DB migrate (why / how / code gate)"


def main():
    doc = Document(DOC_PATH)
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)

    # Refresh Path C command wording if still missing allow env
    for p in doc.paragraphs:
        if p.text.strip() == "npm run deploy:fresh-aws-db  (001–034 → IP_Reset_Core_Sample.js → 035–039)":
            p.text = "IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db  (001–034 → IP_Reset_Core_Sample.js → 035–039)"
        if "Use: Path B = no DB migrate; Path C = npm run deploy:fresh-aws-db" in p.text and "IP_ALLOW" not in p.text:
            p.text = (
                "Use: Path B = no DB migrate (do not set IP_ALLOW_DB_MIGRATE); "
                "Path C = IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db; "
                "existing demo users = IP_ALLOW_DB_MIGRATE=1 npm run db:migrate:sql-only."
            )
        if p.text.strip() == "npm run deploy:fresh-aws-db" and "Path C" in "".join(
            # leave table cells for table loop
            []
        ):
            pass

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t == "npm run deploy:fresh-aws-db":
                    cell.text = "IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db"
                elif t == "npm run db:migrate:sql-only":
                    cell.text = "IP_ALLOW_DB_MIGRATE=1 npm run db:migrate:sql-only"
                elif "Do not run any DB migrate/seed" in t and "IP_ALLOW" not in t:
                    cell.text = "Do not run any DB migrate/seed; do not set IP_ALLOW_DB_MIGRATE"

    if MARKER in joined:
        print("Part S already present — tables/commands refreshed")
        doc.save(DOC_PATH)
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)

    doc.add_heading("Why this issue happened", level=2)
    for line in [
        "Cursor ran a database migration during Path B (app code update only).",
        "Path B must only extract/swap the app, npm install/build, and PM2 restart — RDS must stay untouched.",
        "Agents saw scripts named migrate and ran them; docs alone did not refuse the process.",
        "Separately, failed SQL was sometimes treated as success — now fail-closed with === FAIL === and exit 1.",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_heading("How to stop it (operators + Cursor)", level=2)
    for line in [
        "Path B: never set IP_ALLOW_DB_MIGRATE; never run deploy:fresh-aws-db or db:migrate:*.",
        "Tell Cursor: Path B only — no database.",
        "Path C only when intentional: IP_ALLOW_DB_MIGRATE=1 npm run deploy:fresh-aws-db",
        "SQL only when demo users exist: IP_ALLOW_DB_MIGRATE=1 npm run db:migrate:sql-only",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_heading("How it is stopped in code (current pack)", level=2)
    for line in [
        "scripts/assert-db-migrate-allowed.js refuses migrate/seed unless IP_ALLOW_DB_MIGRATE=1 (or a confirm CLI flag).",
        "Called first by db_exec_sql_file.js, deploy-fresh-aws-db.mjs, db_migrate_sql_only_ip.mjs, IP_Reset_Core_Sample.js, and the handoff runner copies.",
        "Without allow: prints === BLOCKED: database migrate/seed refused === and exits 1.",
        "Casual npm run deploy:fresh-aws-db without the env prefix is refused.",
        "Also see PATH-B-NO-DB-MIGRATE.txt in the handoff zip, AGENTS.md, and .cursor/rules/aws-path-b-no-db.mdc.",
    ]:
        doc.add_paragraph("• " + line)

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH} (Part S added)")


if __name__ == "__main__":
    main()
