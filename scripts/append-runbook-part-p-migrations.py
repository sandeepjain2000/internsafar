#!/usr/bin/env python3
"""Append Part P — migrate-all vs fresh-db + fail-closed migration runners."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
MARKER = "Part P — Migrations: migrate-all vs fresh-db, and fail-closed errors"


def main():
    doc = Document(DOC_PATH)
    existing = "\n".join(p.text for p in doc.paragraphs)
    if MARKER in existing:
        print("Part P already present — skip")
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)

    doc.add_heading("db:migrate:all still exists — it was not removed", level=2)
    doc.add_paragraph(
        "npm run db:migrate:all still runs all IP SQL files (001–039) in manifest order. "
        "Keep it for databases that already have the demo users (or can accept migration 035+)."
    )
    doc.add_paragraph(
        "On a fresh / empty RDS do NOT run migrate-all alone. Migration 035 depends on demo "
        "candidates created by the core seed. Use instead:"
    )
    doc.add_paragraph("• npm run deploy:fresh-aws-db")
    doc.add_paragraph(
        "That command runs: migrations 001–034 → IP_Reset_Core_Sample.js --yes → migrations 035–039."
    )

    doc.add_heading("False “success” when SQL actually failed — fixed", level=2)
    doc.add_paragraph(
        "Earlier Cursor/deploy runs sometimes continued after a migration SQL error and treated "
        "the step as successful. That is not acceptable."
    )
    doc.add_paragraph("Guards in this pack:")
    for line in [
        "db_exec_sql_file.js prints === OK: <file> applied successfully (exit 0) === only after the SQL batch succeeds.",
        "On any PostgreSQL error it prints === FAIL: migration did NOT apply successfully === with code/detail/hint and exits 1.",
        "db:migrate:all and deploy:fresh-aws-db stop at the first failed file (also treat spawn errors and kill signals as failure).",
        "Server NOTICE lines are logged as [pg notice] so they are visible in the terminal.",
        "Never continue past a FAIL banner. Success requires process exit code 0 and OK banners for every file.",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_paragraph(
        "Same guidance is in the handoff README.txt, AWS-DEPLOY.md, and DB-REFERENCE.md."
    )

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
