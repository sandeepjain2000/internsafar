# -*- coding: utf-8 -*-
"""
Sync InternSafar-Test-Cases.xlsx to the live sibling app (2026-09-28).

- Remove rows for retired features (candidate form path, employer manual requests,
  SA Form Registrations / Manual Requests queues, employer Google domain register,
  resume link rows, viral shares, duplicate email-change row)
- Rewrite stale rows (SA signs in on `/`, core accounts, employer register paths,
  candidate profile tabs + required phone, browse tabs, exports, approvals actions)
- Add missing cases (email verify, Hybrid E docs, sticky approval, Suspend/Restore,
  Reset Ethics, ethics lock, Adjust Points, Publish last tab, bulk archive, etc.)
- Recount Index + rewrite Coverage Matrix

  python scripts/sync-internsafar-xlsx-2026-09-28.py            # writes workbook
  python scripts/sync-internsafar-xlsx-2026-09-28.py --dry-run  # residue report only
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib.internsafar_xlsx_cols import COLS, stamp_index_for_results  # noqa: E402

XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
PHASE = "2026-09-28 sync"

BODY = Font(name="Calibri", size=11)
NOTRUN = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(vertical="top", wrap_text=True)
SKIP_SHEETS = {"Index", "Coverage Matrix", "Coverage", "Notes", "How to use", "Meta"}

CAND = "lawsonlclintern+1@gmail.com"
EMP = "placementhubsupport@gmail.com"
EMP_PENDING = "placementhubsupport+3@gmail.com"
SA = "support@placementhub.online"
PRECOND_CORE = (
    f"InternSafar running (sibling internship-portal, npm run dev). Candidate {CAND} / <IP_QA_CORE_PASSWORD>. "
    f"Employer {EMP} / <IP_QA_CORE_PASSWORD>. SuperAdmin {SA} / <IP_QA_CORE_PASSWORD>. All roles sign in on `/` "
    "(#email / #password + captcha)."
)

# Playwright-backed ids after this sync (keep in step with apply-playwright-regression-xlsx.mjs)
AUTOMATED_IDS = {
    "TC-IS-01-001", "TC-IS-01-002", "TC-IS-01-004", "TC-IS-01-005", "TC-IS-01-007",
    "TC-IS-02-001", "TC-IS-02-014", "TC-IS-02-024", "TC-IS-02-025", "TC-IS-02-026",
    "TC-IS-03-005", "TC-IS-03-012", "TC-IS-03-018", "TC-IS-03-024", "TC-IS-03-028", "TC-IS-03-029",
    "TC-IS-04-001", "TC-IS-04-007",
    "TC-IS-06-010",
    "TC-IS-07-007", "TC-IS-07-012", "TC-IS-07-022", "TC-IS-07-023", "TC-IS-07-024", "TC-IS-07-025",
    "TC-IS-09-015", "TC-IS-09-016", "TC-IS-09-017", "TC-IS-09-018",
    "TC-IS-14-012", "TC-IS-14-023", "TC-IS-14-030", "TC-IS-14-031", "TC-IS-14-034",
    "TC-IS-17-008",
    "TC-IS-18-030", "TC-IS-18-039", "TC-IS-18-040", "TC-IS-18-041", "TC-IS-18-042", "TC-IS-18-043",
    "TC-IS-18-044", "TC-IS-18-045", "TC-IS-18-046", "TC-IS-18-047", "TC-IS-18-048",
    "TC-IS-18-049", "TC-IS-18-050", "TC-IS-18-052",
}

# Retired product surfaces — rows deleted outright (history lives in reviews/ audit note).
REMOVE_IDS = {
    "TC-IS-03-004": "Candidate form path returns 410 (see TC-IS-03-028)",
    "TC-IS-03-007": "Form-path referral credit after SA approval — form path retired",
    "TC-IS-03-011": "Employer manual request path returns 410 (see TC-IS-03-029)",
    "TC-IS-03-017": "Form-path captcha — form path retired",
    "TC-IS-03-020": "Employer Google domain register removed (form + email verify now)",
    "TC-IS-03-023": "Employer Google domain register removed",
    "TC-IS-05-004": "Duplicate of TC-IS-06-007 (email change now on /account)",
    "TC-IS-06-009": "Resume links UI removed from candidate profile",
    "TC-IS-14-002": "Form Registrations queue retired (redirect + API 410)",
    "TC-IS-14-003": "Form Registrations queue retired",
    "TC-IS-14-004": "Form Registrations queue retired",
    "TC-IS-14-008": "Manual Requests queue retired",
    "TC-IS-14-009": "Manual Requests queue retired",
    "TC-IS-14-019": "Form Registrations API returns 410 for every role",
    "TC-IS-14-021": "Manual Requests queue retired",
    "TC-IS-18-018": "Viral shares removed (already Obsolete)",
}

# Rows whose expected behaviour changed → reset to Not Run so old Pass is not trusted.
RESET_STATUS = True

EDITS: dict[str, dict[str, str]] = {
    # ---------- 01 Public & Landing ----------
    "TC-IS-01-001": {
        "Description": (
            "1. Open http://localhost:3000/ signed out.\n"
            "2. Confirm Email (#email), Password (#password), captcha, 'Remember this device for 30 days', "
            "Forgot password? and Sign in.\n"
            "3. Confirm 'Create an account' → /register and footer links How it works · Help Center."
        ),
        "Suggestion / Expected Behaviour": (
            "The single login form for all roles is on `/`. Footer shows How it works and Help Center "
            "(Guidelines is not linked from the landing; it stays reachable at /guidelines)."
        ),
    },
    "TC-IS-01-004": {
        "Issue Summary": "How it works / Guidelines / Help load without login",
        "Description": (
            "1. Signed out, open /how-it-works, /guidelines, and /help directly.\n"
            "2. From `/`, click footer links How it works and Help Center.\n"
            "3. Confirm all pages render without login."
        ),
        "Suggestion / Expected Behaviour": (
            "All three routes render as public pages (no 404). Landing footer links reach /how-it-works "
            "and /help. Help Center is content-only (no live ticket API)."
        ),
    },
    "TC-IS-01-009": {
        "Description": (
            "1. Open /register, /register/candidate, and /register/employer at 1280px and 375px.\n"
            "2. /register/candidate: **Sign up with Google** visible, not clipped (no captcha on this page).\n"
            "3. /register/employer: choose **Domain-based** and **Free-email-based** in turn; confirm fields, "
            "captcha, and submit are reachable without overlap."
        ),
        "Suggestion / Expected Behaviour": (
            "Candidate Google button usable at both breakpoints. Employer chooser buttons, form fields, "
            "captcha and submit are clickable. No overlapping footer or clipped CTAs."
        ),
    },
    # ---------- 02 Auth & Access ----------
    "TC-IS-02-001": {
        "Description": (
            f"1. Open `/`.\n2. Candidate: {CAND} / <IP_QA_CORE_PASSWORD> + captcha → Sign in.\n"
            f"3. Sign out. Repeat as employer {EMP} / <IP_QA_CORE_PASSWORD>.\n"
            f"4. Sign out. Repeat as SuperAdmin {SA} / <IP_QA_CORE_PASSWORD> on the same form."
        ),
        "Suggestion / Expected Behaviour": (
            "Candidate → /candidate. Employer → /employer. SuperAdmin → /superadmin. One home form for all roles."
        ),
    },
    "TC-IS-02-006": {
        "Issue Summary": "Employer status decides login: rejected / suspended blocked; pending needs email verify",
        "Description": (
            "1. Sign in as a rejected employer.\n2. Sign in as a suspended employer.\n"
            "3. Sign in as a pending employer whose email is NOT verified.\n"
            "4. Sign in as a pending employer whose email IS verified.\n"
            "5. Sign in as a deactivated user (active=false)."
        ),
        "Suggestion / Expected Behaviour": (
            "Rejected: 'Your employer registration was rejected. Contact support if you believe this is a mistake.' "
            "Suspended: 'Your employer account is suspended. Contact support for help.' "
            "Pending + unverified: 'Verify your email before signing in…' with resend option. "
            "Pending + verified: login succeeds (posting still gated). Deactivated: generic 'Invalid email or password'."
        ),
        "Role(s)": "Employer",
        "Preconditions": "Employer accounts in each approval_status; one pending with and one without email_verified_at.",
    },
    "TC-IS-02-014": {
        "Issue Summary": "SuperAdmin signs in on home `/`; /superadmin/login only redirects",
        "Description": (
            "1. Signed out, open /superadmin/login.\n2. Observe the URL.\n"
            f"3. On `/` sign in with {SA} / <IP_QA_CORE_PASSWORD> + captcha.\n"
            "4. Confirm landing page and SuperAdmin nav."
        ),
        "Suggestion / Expected Behaviour": (
            "/superadmin/login replaces to `/` (no separate SA form, no #sa-email). SuperAdmin credentials on `/` "
            "land on /superadmin with SuperAdmin nav."
        ),
        "Role(s)": "SuperAdmin",
        "Comments / Notes": (
            "Playwright: auth.spec.js 'superadmin login uses standard home form' + 'superadmin signs in and lands on "
            "/superadmin'; regression IS-008 / IS-026. Product reference: src/app/superadmin/login/page.js (redirect)."
        ),
    },
    "TC-IS-02-021": {
        "Suggestion / Expected Behaviour": (
            "Success rows show auth method Password Form (2FA success after OTP). Failures show Unknown account / "
            "Bad Pass / Inactive, plus Employer rejected / Employer suspended where applicable."
        ),
    },
    "TC-IS-02-023": {
        "Description": (
            "1. On `/` (#email/#password + captcha) try: valid candidate; wrong password; unknown email; "
            "pending employer unverified; pending employer verified; suspended employer; 2FA-enabled user; SuperAdmin."
        ),
        "Suggestion / Expected Behaviour": (
            "Candidate → /candidate, employer → /employer, SuperAdmin → /superadmin. Wrong password / unknown → "
            "generic error. Pending unverified → verify-email error. Pending verified → /employer. "
            "Suspended/rejected → status message. 2FA → OTP step before session."
        ),
    },
    # ---------- 03 Registration ----------
    "TC-IS-03-002": {
        "Description": (
            "1. **Google UI:** On /register/candidate, Sign up with Google using a non-Gmail account.\n"
            "2. **API:** POST /api/ip/auth/register-candidate with email person@yahoo.com (no path=form)."
        ),
        "Suggestion / Expected Behaviour": (
            "Non-Gmail rejected before any account is created: 400 'Only Gmail addresses (@gmail.com) are allowed "
            "for candidate registration…'. @googlemail.com still accepted."
        ),
    },
    "TC-IS-03-003": {
        "Description": "1. Sign up with Google using a Gmail that already has an account.",
    },
    "TC-IS-03-008": {
        "Issue Summary": "Domain-based employer register creates pending employer + sends verify email",
        "Description": (
            "1. Open /register/employer → **Domain-based**.\n"
            "2. Fill Company Domain / Website, company name, full name, designation, company domain email, "
            "password (≥8), captcha → submit.\n"
            "3. Confirm success copy and the verification email (no temp password).\n"
            "4. Check SuperAdmin notification / Approvals Pending tab."
        ),
        "Suggestion / Expected Behaviour": (
            "Account created with the employer's own password; approval_status pending; email_verified_at null. "
            "Message: 'Account created. Check your inbox to verify your email, then sign in to upload documents. "
            "Postings unlock after SuperAdmin approval.' SuperAdmin notified. +50 default_signup points."
        ),
        "Legacy ID": "",
    },
    "TC-IS-03-009": {
        "Issue Summary": "Domain path: free / disposable / mismatched email is soft-flagged, not blocked",
        "Description": (
            "1. Domain-based register with website https://example.com and a free mailbox (e.g. @gmail.com).\n"
            "2. Open SuperAdmin Approvals and read the Domain / Risk Tag column for that row."
        ),
        "Suggestion / Expected Behaviour": (
            "Registration succeeds (pending). email_soft_fail true with reasons (free_provider / disposable / "
            "mx_failed); SA notification title 'Employer register — review email'. Approvals shows a risk tag. "
            "No 400 domain-mismatch error."
        ),
    },
    "TC-IS-03-010": {
        "Description": "1. POST /api/ip/auth/register-employer path=domain with all fields + valid captcha but empty website.",
        "Suggestion / Expected Behaviour": (
            "400 'Company domain / website is required for Domain-based registration'. Free-email path does not "
            "require website."
        ),
    },
    "TC-IS-03-012": {
        "Issue Summary": "Employer register missing required fields fails (before captcha)",
        "Description": (
            "1. POST /api/ip/auth/register-employer omitting, one at a time: email, company name, full name, "
            "designation.\n2. Send a bad captcha with otherwise valid fields."
        ),
        "Suggestion / Expected Behaviour": (
            "400 with 'Work email is required' / 'Company name is required' / 'Full name is required' / "
            "'Designation / Role is required'; bad captcha → 'Captcha verification failed'. No user row created."
        ),
        "Role(s)": "Public",
        "Comments / Notes": (
            "Playwright IS-069 (field validation runs before captcha, so no account is created). "
            "Product reference: POST /api/ip/auth/register-employer."
        ),
    },
    "TC-IS-03-013": {
        "Description": "1. Register (Domain or Free-email) with an email that already has an ip_users row.",
        "Suggestion / Expected Behaviour": "409 An account with this email already exists.",
    },
    "TC-IS-03-018": {
        "Issue Summary": "Employer register password BVA: 7 rejected, 8 accepted",
        "Description": (
            "1. POST register-employer with 7-char password (other fields valid).\n"
            "2. Repeat with 8-char password + valid captcha and a unique email.\n"
            "3. Later on /account, change-password still needs upper + digit + special."
        ),
        "Suggestion / Expected Behaviour": (
            "7 → 400 'Password must be at least 8 characters'. 8+ length is enough at register. "
            "Account change-password rules are stricter (TC-IS-02-017)."
        ),
        "Legacy ID": "",
        "Feature": "Password BVA employer register",
        "Preconditions": "None (API); step 2 needs a unique unused email",
        "Comments / Notes": (
            "Playwright IS-069 covers step 1 (7-char → 400, no account). Step 2 creates an account — run manually. "
            "Register rule = length ≥ 8; /account change = ≥ 8 + upper + digit + special."
        ),
    },
    # ---------- 04 Permissions ----------
    "TC-IS-04-002": {
        "Suggestion / Expected Behaviour": (
            "PortalShell compares session role to layout role. Mismatch redirects to `/` for every role "
            "(SuperAdmin layout loginHref is `/` too)."
        ),
    },
    "TC-IS-04-009": {
        "Issue Summary": "/superadmin/login is a redirect-only route",
        "Description": "1. Open /superadmin/login signed out.\n2. Open it signed in as candidate and as SuperAdmin.",
        "Suggestion / Expected Behaviour": (
            "Always replaces to `/` (old bookmarks do not 404). Signed-in users are then routed by role; no "
            "SuperAdmin form is rendered."
        ),
    },
    # ---------- 06 Candidate Profile ----------
    "TC-IS-06-001": {
        "Issue Summary": "Profile complete only when name/phone/college/degree/city/country/resume present",
        "Description": (
            "1. Sign in as candidate on `/`.\n2. Open /candidate/profile.\n"
            "3. Leave one required field blank (name, mobile phone, college, degree, city, country, resume) and save.\n"
            "4. Fill all required fields and save."
        ),
        "Suggestion / Expected Behaviour": (
            "profile_complete stays false until REQUIRED_FOR_COMPLETE is met: name, phone, college, degree, city, "
            "country, resume_url. Blank field error reads '<Label> can't be left blank.' First completion awards +15."
        ),
    },
    "TC-IS-06-002": {
        "Issue Summary": "Six profile tabs render; Privacy & Photo toggles and messaging opt-ins save",
        "Description": (
            "1. Open /candidate/profile.\n2. Confirm tabs: 1. Basics & Contact, 2. Academic, 3. Skills & Experience, "
            "4. Work Readiness, then an 'Optional' divider, then Privacy & Photo and Endorsements (Read-Only) with no number. "
            "On 1. Basics & Contact the header reads 'Setup step 1 of 4'.\n"
            "3. On Privacy & Photo toggle show photo / searchable / hide phone until shortlist; set WhatsApp/Telegram "
            "opt-in.\n4. Save Privacy Settings and reload."
        ),
        "Suggestion / Expected Behaviour": (
            "All six tabs visible. Only the four setup tabs are numbered and show 'Setup step X of 4'; an 'Optional' "
            "divider separates them from Privacy & Photo and Endorsements, which have no number and no step header. "
            "Endorsements read-only. Privacy toggles and opt-ins persist via PUT /api/ip/candidate/profile."
        ),
    },
    "TC-IS-06-004": {
        "Issue Summary": "Candidate profile data export downloads as Excel (.xlsx)",
        "Description": (
            "1. On /candidate/profile click Download Excel (.xlsx).\n"
            "2. GET /api/ip/candidate/export as employer and signed out."
        ),
        "Suggestion / Expected Behaviour": (
            "Candidate gets candidate-portal-export.xlsx. Employer 403, anonymous 401."
        ),
    },
    "TC-IS-06-006": {
        "Description": (
            "1. Open /candidate/profile.\n2. Edit Basics & Contact → Save Basics & Contact.\n"
            "3. Edit Academic rows → Save Academic.\n4. Edit Skills & Experience → save.\n"
            "5. Edit Work Readiness enums → Save Work Readiness.\n6. Privacy: searchable off, hide phone.\n"
            "7. Endorsements tab is read-only."
        ),
    },
    "TC-IS-06-007": {
        "Issue Summary": "Candidate changes login email from Account → Change Email with 6-digit code",
        "Description": (
            "1. Sign in as a gen-accounts candidate (not the shared core unless you restore it).\n"
            "2. Open /account → Change Email; enter a unique unused @gmail.com; request code.\n"
            "3. Submit a wrong code.\n4. Submit the correct code.\n5. Sign in again with the new email."
        ),
        "Suggestion / Expected Behaviour": (
            "Wrong code rejected (invalid or expired). Correct code updates the login email; next login uses it. "
            "Candidate profile page has no separate email-change form."
        ),
        "Feature": "Email change (Account)",
        "Comments / Notes": (
            "Manual helper: npm run qa:manual:email-change (scripts/manual/run-tc-is-06-007-email-change.mjs). "
            "API: /api/ip/candidate/profile/email-change/request + /verify."
        ),
    },
    "TC-IS-06-008": {
        "Issue Summary": "Mobile phone required: blank rejected; invalid number for dial code rejected",
        "Description": (
            "1. Clear phone and Save Basics & Contact.\n2. Set +91 and phone 123 → save.\n3. Set a valid IN mobile → save."
        ),
        "Suggestion / Expected Behaviour": (
            "Blank → 'Mobile phone can't be left blank.' (blocks save / apply unlock). Invalid → validation error from "
            "validateRequiredPhone. Valid number saves."
        ),
    },
    # ---------- 07 Browse ----------
    "TC-IS-07-002": {
        "Description": (
            "1. On /candidate/internships open Filters (Show) and set search, min stipend, work mode, location, "
            "min match / validation, start date.\n2. Change sort.\n3. Wait >450ms and reload."
        ),
        "Suggestion / Expected Behaviour": (
            "Results match GET /api/ip/candidate/internships. Remote listings pass location filter. Last-used "
            "filters/sort restore (tableKey candidate.internships)."
        ),
    },
    "TC-IS-07-003": {
        "Issue Summary": "Browse empty state, List/Cards toggle, mobile layout",
        "Description": (
            "1. Filter until zero results; confirm empty copy.\n2. Toggle List ↔ Cards.\n"
            "3. Resize to ~390px; confirm no horizontal clip and the Filters toggle + saved icon are reachable."
        ),
        "Suggestion / Expected Behaviour": (
            "Empty state explains no matches. View toggle works (List is default). Saved views bar visible. "
            "Mobile layout usable."
        ),
    },
    "TC-IS-07-017": {
        "Issue Summary": "Browse List (default) vs Cards view toggle persists",
        "Description": (
            "1. Fresh browser: open /candidate/internships and note the view.\n2. Switch to Cards; reload.\n"
            "3. Switch back to List."
        ),
        "Suggestion / Expected Behaviour": (
            "Default view is List (table). Choice persists in localStorage key ip_browse_view. Filters/sort unchanged."
        ),
    },
    "TC-IS-07-018": {
        "Issue Summary": "Browse tabs: Unapplied / All Internships / Starting soon / Saved / Recommended for You",
        "Description": (
            "1. Open /candidate/internships.\n2. Click each tab and read its count.\n3. Open /candidate/internships?saved=1."
        ),
        "Suggestion / Expected Behaviour": (
            "Each tab filters consistently with its count (Unapplied/Starting soon send chip=, Saved sends "
            "savedOnly=1). ?saved=1 opens the Saved tab. Tab choice sticks for the browser session."
        ),
    },
    # ---------- 09 Employer pipeline ----------
    "TC-IS-09-001": {
        "Suggestion / Expected Behaviour": (
            "Pipeline statuses: applied, shortlisted, interviewing, rejected, hired (completed via completions). "
            "Interviewing requires date/time; invalid Meet URL rejected. Candidate gets 'Interview invitation "
            "received' with a Messages link; application shows Interview Scheduled."
        ),
    },
    "TC-IS-09-017": {
        "Issue Summary": "Employer dashboard Action center shows one priority action with count and link",
        "Description": (
            "1. Sign in as core employer.\n2. Open `/employer`.\n3. Read the Action center card and follow it.\n"
            "4. Repeat as a pending employer with no documents."
        ),
        "Suggestion / Expected Behaviour": (
            "Card (data-testid=employer-action-center) shows a number and one hint, first match wins: Upload "
            "verification documents → Finish your company profile → Applications waiting for your review (3+ days) "
            "/ No applications waiting for review → Waiting for Final Approval. Links to /employer/profile or "
            "/employer/internships."
        ),
        "Comments / Notes": "Playwright IS-063 / JOURNEY-EMP-01. Redesigned in ba5a08e (single priority card).",
    },
    "TC-IS-09-011": {
        "Issue Summary": "Applicant export: Excel only, or ZIP with applicants.xlsx + resumes/",
        "Description": (
            "1. On /employer/internships/[id] select applicants → export without CVs.\n"
            "2. Repeat with CVs included.\n3. Open both files."
        ),
        "Suggestion / Expected Behaviour": (
            "Without CVs → applicants-export.xlsx. With CVs → ZIP containing applicants.xlsx plus resumes/. "
            "Hidden phones omitted per privacy rules."
        ),
    },
    "TC-IS-09-014": {
        "Description": (
            "1. Open /employer/internships/new.\n2. Walk tabs Details → Schedule → Hours & Engagement → "
            "Compensation → Eligibility → Screening (Next Tab).\n3. Save Draft.\n4. Re-open edit and confirm values."
        ),
        "Suggestion / Expected Behaviour": (
            "Draft persists without spending points. Body sections, locations, stipend, eligibility and screening "
            "questions round-trip. Publish costs 50 pts when used."
        ),
    },
    # ---------- 12 / 13 ----------
    "TC-IS-13-002": {
        "Description": (
            "1. Track ledger after: candidate Google signup, employer register, complete profile, first apply, "
            "employer publish, successful referral."
        ),
        "Suggestion / Expected Behaviour": (
            "default_signup +50 (candidate and employer). profile_complete +15 once. application_spend -5. "
            "first_application_bonus +10 once. posting_spend -50. referral_bonus +25. Convert-credits retired."
        ),
    },
    "TC-IS-13-004": {
        "Suggestion / Expected Behaviour": (
            "Chronological ledger sum equals balance (includes sa_manual_credit / sa_manual_debit). Referral +25 is "
            "credited at registration for both candidate (Google) and employer referrals."
        ),
    },
    # ---------- 14 SuperAdmin ----------
    "TC-IS-14-001": {
        "Description": "1. Open /superadmin as SuperAdmin.\n2. Compare pending counts with Employer approvals and Documents.",
        "Suggestion / Expected Behaviour": (
            "Stats load for SuperAdmin only; counts match the live queues (no Form Registrations / Manual Requests). "
            "Candidate hitting /superadmin is bounced."
        ),
    },
    "TC-IS-14-005": {
        "Description": (
            "1. Register a Domain or Free-email employer and verify the email.\n"
            "2. Employer uploads a document; SuperAdmin approves it in Documents.\n"
            "3. On /superadmin/approvals Pending tab, Approve.\n"
            "4. Employer Accept & Saves ethics, completes profile, publishes."
        ),
        "Suggestion / Expected Behaviour": (
            "approval_status approved ('Final Employer Approval Granted' notice + email). Posting unlocks once email "
            "verified + ethics saved + profile complete."
        ),
    },
    "TC-IS-14-006": {
        "Issue Summary": "Approvals tabs: Pending Approve/Reject; Approved Suspend/Reject; Suspended Restore/Reject",
        "Description": (
            "1. Pending tab: Reject one row with a preset reason.\n2. Approved tab: Suspend (confirm dialog).\n"
            "3. Suspended tab: Restore.\n4. Approved tab: Reject with reason."
        ),
        "Suggestion / Expected Behaviour": (
            "Rows move to the matching tab. Reject stores rejection_reason (required for row reject). Suspended and "
            "rejected employers cannot sign in or post. Restore sets approved again without re-running the document check."
        ),
    },
    "TC-IS-14-007": {
        "Suggestion / Expected Behaviour": "400 'approvalStatus Must Be One Of approved, rejected, suspended, pending'.",
    },
    "TC-IS-14-010": {
        "Issue Summary": "SuperAdmin Documents lists active employer uploads and approves/rejects",
        "Description": (
            "1. Employer uploads Shop Act and Other (with name).\n2. SuperAdmin opens /superadmin/documents.\n"
            "3. Open a file; approve one, reject one."
        ),
        "Suggestion / Expected Behaviour": (
            "Only active (non-superseded) docs appear. Type label shows 'Other - <name>'. Status badges Title Case "
            "(Pending / Approved / Rejected). Employer notified on review."
        ),
    },
    "TC-IS-14-022": {
        "Description": f"1. Sign in on `/` as {SA}.\n2. Walk the left nav.\n3. Confirm there is no Notifications item.",
    },
    # ---------- 18 Cross-cutting / employer ----------
    "TC-IS-18-001": {
        "Description": (
            "1. As employer save profile missing HQ City or Contact Phone.\n"
            "2. Fill company name, website, work email, industry, HQ city, contact name, contact phone, business "
            "entity type.\n3. Accept & Save all Guidelines & Ethics items."
        ),
        "Suggestion / Expected Behaviour": (
            "profile_complete only after REQUIRED_FOR_COMPLETE fields + ethics saved. Until then posting returns 403 "
            "'Complete your employer profile before posting.' (or the ethics message if ethics unsaved)."
        ),
    },
    "TC-IS-18-002": {
        "Issue Summary": "Upload verification documents by type (Shop Act / LLP registration / Business PAN / Other)",
        "Description": (
            "1. Profile & docs → Verification Documents.\n2. Upload one file per type; for Other give a name.\n"
            "3. Try Other without a name.\n4. Try an unsupported file type."
        ),
        "Suggestion / Expected Behaviour": (
            "Each upload stored as pending for its type. Other without name → 'Document name is required when type is "
            "Other'. Bad type/size errors cleanly."
        ),
    },
    "TC-IS-18-003": {
        "Issue Summary": "Guidelines & Ethics: Accept/Reject per item, Save Acknowledgements locks",
        "Description": (
            "1. Profile & docs → Guidelines & Ethics.\n2. Reject one item and save.\n3. Accept all and Save "
            "Acknowledgements.\n4. Try to change any item."
        ),
        "Suggestion / Expected Behaviour": (
            "Each item has Accept / Reject (not checkboxes). Any Reject keeps profile incomplete. All Accept + save "
            "stamps ethics_accepted_at and locks the controls."
        ),
    },
    "TC-IS-18-004": {
        "Description": "1. View /employer as pending (email verified) vs approved employer.\n2. Click post internship.",
        "Suggestion / Expected Behaviour": (
            "Pending: can sign in and upload docs; posting blocked with 'must be approved by SuperAdmin' message. "
            "Approved + verified + ethics saved + profile complete: /employer/internships/new works."
        ),
    },
    "TC-IS-18-019": {
        "Suggestion / Expected Behaviour": (
            "?ref= survives to /register/employer. Referrer gets +25 referral_bonus at registration (either path) "
            "and a 'Referral bonus earned' notice."
        ),
    },
    "TC-IS-18-027": {
        "Description": (
            "1. Register Google-path candidate.\n2. Register an employer (verify link).\n3. Enable 2FA.\n"
            "4. Request password reset.\n5. Check the QA override inbox."
        ),
        "Suggestion / Expected Behaviour": (
            "With OUTBOUND_EMAIL_OVERRIDE set, candidate temp password, employer verify link, OTPs and reset mail "
            "land in the override inbox. Registration still succeeds if mail fails (warning shown)."
        ),
    },
    "TC-IS-18-028": {
        "Description": (
            "1. Candidate: Notifications → /candidate/notifications.\n2. Employer: Notifications → /employer/notifications.\n"
            "3. SuperAdmin: Dashboard, Employer approvals, Documents, Postings, Posting Share Rewards, Adjust Points, "
            "Login report, Listing reports, Messages, Feature ideas, Account — no Notifications.\n4. Mobile drawer."
        ),
        "Suggestion / Expected Behaviour": "Labels match src/lib/ipNav.js. SuperAdmin has no notifications page.",
    },
    "TC-IS-18-029": {
        "Issue Summary": "Sign out returns every role to landing `/`",
        "Description": "1. Sign out from candidate, employer, and SuperAdmin shells.",
        "Suggestion / Expected Behaviour": "All roles land on `/` (loginHref is `/` for every layout).",
    },
    "TC-IS-18-032": {
        "Issue Summary": "Captcha on employer register, forgot-password, and login",
        "Description": (
            "1. /register/employer (Domain and Free-email): wrong captcha → fail closed.\n"
            "2. Forgot-password: wrong captcha → fail closed.\n3. Home login: wrong captcha → fail closed.\n"
            "4. Candidate Google register has no captcha field."
        ),
        "Suggestion / Expected Behaviour": (
            "Captcha enforced on employer register, forgot-password and login. Candidate Google register is exempt."
        ),
    },
    "TC-IS-18-034": {
        "Description": (
            f"1. POST /api/ip/bootstrap (or load `/`).\n2. Sign in {SA} on `/`.\n"
            f"3. Confirm {CAND} and {EMP} still exist."
        ),
    },
}

NEW_CASES: list[dict[str, str]] = []


def case(tc_id, sheet, feature, typ, summary, steps, expected, roles, pre, auto="Manual", sev="P1", notes=""):
    NEW_CASES.append({
        "ID": tc_id,
        "sheet": sheet,
        "Module / Section": sheet,
        "Feature": feature,
        "Type": typ,
        "Issue Summary": summary,
        "Description": steps,
        "Suggestion / Expected Behaviour": expected,
        "Role(s)": roles,
        "Preconditions": pre,
        "Automation": auto,
        "Severity / Priority": sev,
        "Doc Status": "Ready for QA",
        "Phase": PHASE,
        "Test Status": "Not Run",
        "Legacy ID": "",
        "Comments / Notes": notes,
    })


S02, S03, S07, S09 = "02 Auth & Access", "03 Registration", "07 Browse Save Apply", "09 Employer Postings Pipeline"
S12, S13, S14, S17, S18 = "12 Messages", "13 Points & Referrals", "14 SuperAdmin Ops", "17 Files & Uploads", "18 Cross-cutting"

case("TC-IS-02-028", S02, "Employer email verify gate", "Negative",
     "Pending employer with unverified email cannot sign in; resend offered",
     "1. Register a new employer; do not click the verify link.\n2. Sign in on `/` with that email/password.\n3. Use the resend option.",
     "Login blocked: 'Verify your email before signing in. Check your inbox, or resend the verification email below.' Resend sends a new link (see TC-IS-03-027).",
     "Employer", "Freshly registered employer (email_verified_at null)", sev="P0")
case("TC-IS-02-029", S02, "Employer email verify gate", "Functional",
     "Pending employer with verified email can sign in; posting stays gated",
     "1. Verify the new employer's email.\n2. Sign in on `/`.\n3. Open Profile & docs and upload a document.\n4. Try /employer/internships/new (create or draft).",
     "Login succeeds → /employer. Documents upload works. Posting returns 403 'Your employer account must be approved by SuperAdmin before posting'.",
     "Employer", f"Pending verified employer (e.g. {EMP_PENDING} if verified)", sev="P0",
     notes="Chicken-and-egg fix 2026-09-23: login allowed after email verify so docs can be uploaded before Final Approval.")

case("TC-IS-03-024", S03, "Employer register chooser", "UI",
     "Employer register offers Domain-based and Free-email-based paths (no Google)",
     "1. Open /register/employer signed out.\n2. Read the chooser.\n3. Open each path.",
     "Heading 'Employer Registration' with Domain-based and Free-email-based buttons. No Google button. Domain path shows Company Domain / Website + Company Domain Email; Free-email path shows Email and no website field. Both show password + captcha.",
     "Public", "None", auto="Automated", sev="P1", notes="Playwright IS-070")
case("TC-IS-03-025", S03, "Employer register (free email)", "Functional",
     "Free-email employer register creates pending employer without website",
     "1. /register/employer → Free-email-based.\n2. Fill company, full name, designation, a personal mailbox, password ≥8, captcha → submit.",
     "Account created pending (registration_source free_email), verify email sent, no temp password. Appears in Approvals Pending with Free-email path tag.",
     "Public, SuperAdmin", "Unique unused mailbox", sev="P0")
case("TC-IS-03-026", S03, "Employer email verify", "Functional",
     "Employer verify link marks email verified; reused/expired link fails cleanly",
     "1. Open the verify link from the register email.\n2. Open the same link again.\n3. Open a tampered link.",
     "First open → verified page; email_verified_at set; login now allowed. Reuse/tampered → friendly invalid/expired state, no crash.",
     "Employer", "Freshly registered employer", sev="P0")
case("TC-IS-03-027", S03, "Employer email verify resend", "Edge",
     "Resend verification respects ~45s cooldown and does not reveal account existence",
     "1. On the register success screen click resend.\n2. Click again immediately.\n3. POST /api/ip/auth/employer-email-verify/resend with an unknown email.",
     "First → 'If that email needs verification, a new link has been sent.' with cooldown ~45s. Immediate retry → 429 with retryAfterSec; button disabled with countdown. Unknown email → same generic ok.",
     "Public", "Freshly registered employer", sev="P1")
case("TC-IS-03-028", S03, "Retired candidate form path", "Regression",
     "Candidate form-path registration API returns 410",
     "1. POST /api/ip/auth/register-candidate with {path:'form', email, password, captcha}.",
     "410 'Form registration is no longer available. Continue with Google on the candidate register page.' No user created.",
     "Public", "None", auto="Automated", sev="P1", notes="Playwright IS-067")
case("TC-IS-03-029", S03, "Retired employer manual request", "Regression",
     "Employer manualRequest registration returns 410",
     "1. POST /api/ip/auth/register-employer with {manualRequest:true, …}.",
     "410 'Manual employer requests are no longer accepted. Register with Domain or Free-email on /register/employer.'",
     "Public", "None", auto="Automated", sev="P1", notes="Playwright IS-068")
case("TC-IS-03-030", S03, "Employer register password", "Functional",
     "Employer signs in with the password chosen at register (no temp password email)",
     "1. Register an employer with password P.\n2. Verify email.\n3. Sign in on `/` with P.\n4. Check inbox for any temp-password mail.",
     "Login works with P. Only the verify email and 'registration received' ack are sent — no temp password.",
     "Employer", "Freshly registered employer", sev="P0")

case("TC-IS-07-025", S07, "Browse defaults", "Regression",
     "Browse opens on Unapplied tab by default",
     "1. Fresh browser session: open /candidate/internships.\n2. Read which tab is selected.",
     "Unapplied tab has aria-selected=true. Default view is List when the browser has no saved view.",
     "Candidate", "Core candidate", auto="Automated", sev="P2", notes="Playwright IS-078")

case("TC-IS-09-018", S09, "Posting form tabs", "Regression",
     "Publish Now appears only on the last tab (Screening)",
     "1. Open /employer/internships/new.\n2. On Details, look for Publish Now vs Next Tab.\n3. Open Screening.",
     "Details..Eligibility show Next Tab and Save Draft, no Publish Now. Screening shows Publish Now. Header has Back to Postings.",
     "Employer", "Approved core employer", auto="Automated", sev="P1", notes="Playwright IS-076")
case("TC-IS-09-019", S09, "Stipend range", "Negative",
     "Stipend max below min is rejected",
     "1. New posting → Compensation: Fixed Stipend, min 10000, max 5000.\n2. Save Draft.\n3. Set max 15000 and save.",
     "Step 2 error 'Stipend maximum must be greater than or equal to the minimum.' Step 3 saves. Incentive-Based hides the range check.",
     "Employer", "Approved core employer", sev="P2")
case("TC-IS-09-020", S09, "Applicant CV download", "Functional",
     "Single candidate download from full profile returns ZIP (candidate.xlsx + resume)",
     "1. /employer/candidates/{id} for an applicant → Download.\n2. Open the ZIP.",
     "candidate-export.zip with candidate.xlsx (only fields the employer may see) and resumes/ when a CV exists. Phone only when shortlist/privacy allows.",
     "Employer", "Candidate who applied to the employer's posting", sev="P2")
case("TC-IS-09-021", S09, "Employer exports", "Functional",
     "Employer dashboard overview and profile export download as .xlsx",
     "1. /employer → Export overview (.xlsx).\n2. /employer/profile → Download Excel (.xlsx).",
     "Both download valid .xlsx files (employer-export-<company>.xlsx, employer-profile-<company>.xlsx). Candidate gets 403 on these APIs.",
     "Employer", "Core employer", sev="P2")

case("TC-IS-12-011", S12, "Bulk archive", "Functional",
     "Messages Select All + Archive Selected / Unarchive Selected",
     "1. Open /employer/messages (repeat /candidate/messages).\n2. Select two threads → Archive Selected.\n3. Archived view → Unarchive Selected.",
     "Selected threads move between views for this role only; other side's inbox unchanged. Counts on buttons match selection.",
     "Candidate, Employer", "Two or more threads", sev="P2")
case("TC-IS-12-012", S12, "Loading vs empty", "UI",
     "Notifications show Please Wait… while loading, not the empty state",
     "1. Hard-refresh /employer/notifications and /candidate/notifications (throttle network if needed).",
     "Spinner / 'Please Wait…' until data arrives; empty 'all caught up' copy only when the list is truly empty.",
     "Candidate, Employer", "Core accounts", sev="P2")

case("TC-IS-13-005", S13, "SA manual points", "Functional",
     "SuperAdmin point adjustment shows in user ledger and in-app notice",
     "1. SA adds 10 points to a candidate via Adjust Points.\n2. Candidate opens Refer & earn ledger and Notifications.",
     "Ledger row sa_manual_credit +10; balance +10. Notification 'Points Added' — 'Admin has added 10 points to your account.' No email.",
     "SuperAdmin, Candidate", "Core accounts", sev="P1")

case("TC-IS-14-024", S14, "Final Approval document gate", "Negative",
     "Final Approval blocked until ≥1 approved and 0 pending active documents",
     "1. Approve an employer with no documents.\n2. With one pending document.\n3. With only rejected documents.\n4. With 1 approved + 1 rejected + 0 pending.",
     "1 → 'has not uploaded any verification documents yet…'. 2 → 'approve this employer's pending document…'. 3 → 'approve at least one verification document…'. 4 → approval succeeds.",
     "SuperAdmin", "Pending employers in each document state", sev="P0")
case("TC-IS-14-025", S14, "Sticky approval", "Functional",
     "After Final Approval a new pending document does not block posting",
     "1. Approved employer uploads a replacement document.\n2. Check Documents (Pending).\n3. Employer publishes a posting.",
     "New doc is Pending in Documents; approval_status stays approved; posting still works. Only an explicit SA Suspend/Reject/Pending stops posting.",
     "Employer, SuperAdmin", "Approved employer able to post", sev="P0")
case("TC-IS-14-026", S14, "Suspend / Restore", "Functional",
     "Suspend approved employer blocks login; Restore re-enables even with a pending document",
     "1. As the approved employer upload a new document and leave it Pending.\n2. Approved tab → Suspend (confirm).\n3. Employer tries to sign in.\n4. Suspended tab → Restore.\n5. Try PATCH approvalStatus=suspended on a pending or rejected employer.",
     "Suspend → 'Employer Account Suspended' notice; login shows suspended message. Restore → approved, 'Employer Account Restored' notice, login and posting work — a pending document does not block Restore (sticky approval). Suspending a non-approved employer returns 400 'Only approved employers can be suspended.'",
     "SuperAdmin, Employer", "Approved employer", sev="P0",
     notes="Restore (suspended → approved) skips the Final Approval document check; suspend is only allowed from approved.")
case("TC-IS-14-027", S14, "Reject approved employer", "Functional",
     "Reject approved or suspended employer requires a reason and blocks login",
     "1. Approved tab → row Reject; submit with empty custom reason (preset Other).\n2. Pick a preset reason and submit.\n3. Employer tries to sign in.",
     "Empty reason → 'Rejection reason is required'. With reason → Rejected tab, rejection_reason stored, email + notice sent. Login shows rejected message.",
     "SuperAdmin, Employer", "Approved employer", sev="P0")
case("TC-IS-14-028", S14, "Reset Ethics", "Functional",
     "Reset Ethics unlocks acknowledgements and blocks posting until re-saved",
     "1. Approvals → select employer → Reset Ethics.\n2. Employer tries to create a draft.\n3. Employer Accept & Saves all items.\n4. Create a draft again.",
     "Reset clears ethics_acks / ethics_accepted_at and sets profile_complete false. Step 2 → 403 'Accept all Guidelines & Ethics acknowledgements and save them before posting.' Step 4 succeeds.",
     "SuperAdmin, Employer", "Approved employer with locked ethics", sev="P0")
case("TC-IS-14-029", S14, "Adjust Points", "Functional",
     "Adjust Points: search user, add or deduct with optional note",
     "1. Open /superadmin/points.\n2. Search a candidate by email → Adjust → Add 25 with no note → confirm.\n3. Adjust → Deduct 5 with a note.",
     "Balance updates (+25 then -5), toast confirms. Ledger reasons sa_manual_credit / sa_manual_debit with note in meta. User gets in-app notice only.",
     "SuperAdmin", "Core accounts", sev="P1")
case("TC-IS-14-030", S14, "Adjust Points guards", "Negative",
     "Adjust Points rejects zero/non-integer, below-zero deduct, SA targets, and non-SA callers",
     "1. Open /superadmin/points (page loads).\n2. POST /api/ip/superadmin/points delta 0 and delta 1.5.\n3. Deduct more than the balance.\n4. Target a SuperAdmin user.\n5. POST as candidate.",
     "2 → 400 'delta must be a non-zero integer'. 3 → 400 'Cannot deduct N points — balance is B (would go below 0)'. 4 → 400 'Only candidate or employer accounts can be adjusted'. 5 → 403. No per-action cap.",
     "SuperAdmin, Candidate", "Core accounts", auto="Automated", sev="P1", notes="Playwright IS-072 (steps 1, 2, 5)")
case("TC-IS-14-031", S14, "Retired SA queues", "Regression",
     "Form Registrations / Manual Requests / Viral pages redirect; their APIs return 410",
     "1. As SA open /superadmin/form-registrations and /superadmin/requests.\n2. Open /superadmin/viral.\n3. GET /api/ip/superadmin/form-registrations and /api/ip/superadmin/requests.",
     "1 → /superadmin/approvals. 2 → /superadmin. 3 → 410 with 'retired' message.",
     "SuperAdmin", "SuperAdmin session", auto="Automated", sev="P2", notes="Playwright IS-071")
case("TC-IS-14-032", S14, "Listing reports", "Functional",
     "Listing reports queue shows candidate reports and updates status",
     "1. Candidate reports a listing from detail (Report).\n2. SA opens /superadmin/listing-reports.\n3. Mark reviewed / dismissed.",
     "Report appears under Open. PATCH status open|reviewed|dismissed moves it; invalid status → 400.",
     "Candidate, SuperAdmin", "Published internship", sev="P2")
case("TC-IS-14-033", S14, "Delete employer", "Functional",
     "Delete Selected soft-deletes employer (rejected + inactive)",
     "1. Approvals → select a test employer → Delete Selected.\n2. Employer tries to sign in.",
     "Row moves to Rejected with reason 'Deleted By SuperAdmin'; user active=false; login fails with generic error.",
     "SuperAdmin", "Disposable test employer", sev="P2")
case("TC-IS-14-034", S14, "Approvals UI", "Regression",
     "Approvals page shows Pending / Approved / Suspended / Rejected tabs with tab-specific bulk actions",
     "1. Open /superadmin/approvals.\n2. Switch each tab and read the header buttons.",
     "Pending: Approve Selected + Reject Selected. Approved: Suspend Selected + Reject Selected. Suspended: Restore Selected + Reject Selected. Delete Selected and Reset Ethics on all.",
     "SuperAdmin", "SuperAdmin session", auto="Automated", sev="P1", notes="Playwright IS-073")

case("TC-IS-17-006", S17, "Document slots (Hybrid E)", "Functional",
     "Re-upload of the same document type supersedes the previous file and goes Pending",
     "1. Employer uploads Business PAN (approved later by SA).\n2. Upload Business PAN again.\n3. Check Profile & docs and SA Documents.",
     "Only one active Business PAN; the old file is kept as history (superseded) and hidden from queues. New file is Pending. Max four active docs (one per type).",
     "Employer, SuperAdmin", "Employer with an approved document", sev="P0")
case("TC-IS-17-007", S17, "Logo URL", "UI",
     "Uploaded logo URL is hidden; Change via logo URL reveals an empty field",
     "1. Profile & docs → Company Details: click the logo frame and upload an image.\n2. Look for any URL text box.\n3. Click Change via logo URL.",
     "Image fills the frame; no storage URL is shown. URL field appears empty only after Change via logo URL.",
     "Employer", "Core employer", sev="P2")
case("TC-IS-17-008", S17, "File access control", "Security",
     "Candidate cannot fetch another user's stored object via /api/ip/files",
     "1. As candidate GET /api/ip/files?key=internship-portal/candidates/<other user>/resume/x.pdf.",
     "403 Forbidden. No bytes returned.",
     "Candidate", "Core candidate", auto="Automated", sev="P0", notes="Playwright IS-057")

case("TC-IS-18-049", S18, "Employer profile tabs", "Regression",
     "Employer profile shows five tabs; Contact & Location order Country → State → City → Phone → Work Email",
     "1. Open /employer/profile.\n2. Read tabs.\n3. Open Contact & Location and read field order.\n4. Change HQ Country.",
     "Tabs: Company Details, Contact & Location, About & Visibility, Guidelines & Ethics, Verification Documents (one section at a time). Order HQ Country, HQ State / Province, HQ City, Contact Phone, Work Email. Country is single-select; changing it clears State/City.",
     "Employer", "Core employer", auto="Automated", sev="P1", notes="Playwright IS-074 (tabs + order)")
case("TC-IS-18-050", S18, "Ethics lock", "Security",
     "Locked ethics cannot be revoked via API",
     "1. Employer with saved ethics PUT /api/ip/employer/profile {ethics_acks:{}}.",
     "403 'Guidelines & Ethics are locked after save. Contact SuperAdmin to reset acknowledgements.' Stored acks unchanged.",
     "Employer", "Core employer with locked ethics", auto="Automated", sev="P0", notes="Playwright IS-075 (skips if not locked)")
case("TC-IS-18-051", S18, "Posting gate order", "Negative",
     "Posting gate checks approval → email verified → ethics saved → profile complete",
     "1. POST /api/ip/employer/internships as employers failing each gate in turn.",
     "403 messages in order: 'must be approved by SuperAdmin before posting' → 'Verify your email before posting…' → 'Accept all Guidelines & Ethics acknowledgements and save them before posting.' → 'Complete your employer profile before posting.' Drafts are gated too.",
     "Employer", "Employers in each gate state", sev="P0")
case("TC-IS-18-052", S18, "Filters toggle", "Regression",
     "List Filters button toggles Show / Hide and opens the panel",
     "1. /candidate/internships: read the Filters button.\n2. Click it.\n3. Click again.",
     "Button reads 'Filters Show' when closed and 'Filters Hide' when open (aria-expanded matches). Panel .ip-tf__panel opens/closes.",
     "Candidate, Employer", "Core accounts", auto="Automated", sev="P2", notes="Playwright IS-077")
case("TC-IS-18-053", S18, "Title Case badges", "UI",
     "Status badges use Title Case across lists",
     "1. Check approval badge on employer profile, document badges, applicant statuses, SA documents/postings, candidate applications.",
     "Badges read e.g. Approved / Pending / Interview Scheduled / Closing Soon — never raw lower-case enum values.",
     "All", "Core accounts", sev="P2")

# Stale phrases that should not survive (reported after run).
RESIDUE_PATTERNS = [
    r"sa-email", r"sa-password", r"/superadmin/login with", r"Form registrations?", r"form-registrations",
    r"Manual requests?", r"manual_request", r"shreekar", r"Filters On/Off", r"Continue with Google",
    r"temp password emailed", r"Guidelines · Help", r"Cards vs List", r"quick chips", r"path=form",
    r"Form path", r"five tabs",
]

GLOBAL_SUBS = [
    (re.compile(r"(SA|SuperAdmin)(\s*:?\s*)placementhubsupport@gmail\.com"), rf"\1\2{SA}"),
    (re.compile(r"shreekar\.nyayapathi23\+2@vit\.edu"), EMP),
    (re.compile(r"SuperAdmin on /superadmin/login \(#sa-email, #sa-password\)"),
     "SuperAdmin also on `/` (#email, #password); /superadmin/login only redirects to `/`"),
    (re.compile(r"SuperAdmin: only /superadmin/login with #sa-email/#sa-password\.?"),
     "SuperAdmin signs in on `/` too."),
    (re.compile(r"quick chips"), "browse tabs"),
    (re.compile(r"auth\.js active / form_approval_status"), "auth.js employer approval_status + email_verified_at"),
]


def find_header(ws):
    for r in range(1, 12):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 40) + 1)]
        if vals and vals[0] in ("ID", "TC ID"):
            return r, {str(v): i + 1 for i, v in enumerate(vals) if v}
    return None, {}


def style_cell(cell, fill=None):
    cell.font = BODY
    cell.alignment = WRAP
    if fill is not None:
        cell.fill = fill


def reset_status(ws, r, cols):
    st = cols.get("Test Status")
    if st:
        ws.cell(r, st).value = "Not Run"
        style_cell(ws.cell(r, st), NOTRUN)
    for key in ("Actual Result", "Date Verified"):
        c = cols.get(key)
        if c:
            ws.cell(r, c).value = None


def sev_bucket(v):
    s = str(v or "").strip()
    if s in ("P0", "Critical"):
        return "P0"
    if s in ("P1", "High"):
        return "P1"
    return "P2+"


def main(dry_run=False):
    wb = load_workbook(XLSX)
    stats = Counter()
    residue = []

    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        hr, cols = find_header(ws)
        if not hr:
            continue
        id_col = cols["ID"]

        for r in range(ws.max_row, hr, -1):
            tid = str(ws.cell(r, id_col).value or "").strip()
            if tid in REMOVE_IDS:
                ws.delete_rows(r)
                stats["removed"] += 1

        for r in range(hr + 1, ws.max_row + 1):
            tid = str(ws.cell(r, id_col).value or "").strip()
            if not tid.startswith("TC-IS-"):
                continue
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if not isinstance(v, str):
                    continue
                nv = v
                for rx, rep in GLOBAL_SUBS:
                    nv = rx.sub(rep, nv)
                if nv != v:
                    ws.cell(r, c).value = nv
                    stats["text_subs"] += 1
            pre = cols.get("Preconditions")
            if pre and "InternSafar running" in str(ws.cell(r, pre).value or ""):
                ws.cell(r, pre).value = PRECOND_CORE
            if tid in EDITS:
                changed = False
                for key, val in EDITS[tid].items():
                    c = cols.get(key)
                    if c and ws.cell(r, c).value != val:
                        ws.cell(r, c).value = val
                        style_cell(ws.cell(r, c))
                        changed = True
                if changed:
                    if RESET_STATUS:
                        reset_status(ws, r, cols)
                    ph = cols.get("Phase")
                    if ph:
                        ws.cell(r, ph).value = PHASE
                    stats["edited"] += 1
            sv = cols.get("Severity / Priority")
            if sv and str(ws.cell(r, sv).value or "") in ("High", "Medium"):
                ws.cell(r, sv).value = "P1" if ws.cell(r, sv).value == "High" else "P2"
            ac = cols.get("Automation")
            if ac:
                want = "Automated" if tid in AUTOMATED_IDS else "Manual"
                if str(ws.cell(r, ac).value or "") != want:
                    ws.cell(r, ac).value = want
                    stats["automation_changed"] += 1

    sheets = {ws.title: ws for ws in wb.worksheets}
    for rec in NEW_CASES:
        ws = sheets[rec["sheet"]]
        hr, cols = find_header(ws)
        existing = {str(ws.cell(r, cols["ID"]).value or "") for r in range(hr + 1, ws.max_row + 1)}
        if rec["ID"] in existing:
            for r in range(hr + 1, ws.max_row + 1):
                if str(ws.cell(r, cols["ID"]).value or "") == rec["ID"]:
                    for key, val in rec.items():
                        if key in ("sheet", "Test Status") or key not in cols:
                            continue
                        ws.cell(r, cols[key]).value = val
                        style_cell(ws.cell(r, cols[key]))
            stats["new_updated"] += 1
            continue
        r = ws.max_row + 1
        for key in COLS:
            c = cols.get(key)
            if c:
                ws.cell(r, c).value = rec.get(key)
                style_cell(ws.cell(r, c), NOTRUN if key == "Test Status" else None)
        stats["added"] += 1

    # Sort case rows by ID inside each sheet so new cases sit in order.
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        hr, cols = find_header(ws)
        if not hr:
            continue
        rows = []
        for r in range(hr + 1, ws.max_row + 1):
            vals = [ws.cell(r, c) for c in range(1, ws.max_column + 1)]
            if not str(vals[0].value or "").startswith("TC-IS-"):
                continue
            rows.append([(c.value, c._style) for c in vals])
        rows.sort(key=lambda row: str(row[0][0]))
        for i, row in enumerate(rows):
            for j, (val, sty) in enumerate(row):
                cell = ws.cell(hr + 1 + i, j + 1)
                cell.value = val
                cell._style = sty

    # Residue scan
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        hr, cols = find_header(ws)
        if not hr:
            continue
        for r in range(hr + 1, ws.max_row + 1):
            tid = str(ws.cell(r, cols["ID"]).value or "")
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if not isinstance(v, str):
                    continue
                for pat in RESIDUE_PATTERNS:
                    if re.search(pat, v, re.I):
                        residue.append(f"{tid} [{ws.cell(hr, c).value}] ~ {pat}")

    # Index recount
    total = Counter()
    if "Index" in wb.sheetnames:
        idx = wb["Index"]
        for r in range(5, 24):
            name = idx.cell(r, 2).value
            if name not in sheets or name in SKIP_SHEETS:
                continue
            ws = sheets[name]
            hr, cols = find_header(ws)
            cnt = Counter()
            for rr in range(hr + 1, ws.max_row + 1):
                if str(ws.cell(rr, cols["ID"]).value or "").startswith("TC-IS-"):
                    cnt["n"] += 1
                    cnt[sev_bucket(ws.cell(rr, cols["Severity / Priority"]).value)] += 1
            idx.cell(r, 3).value = str(cnt["n"])
            idx.cell(r, 4).value = str(cnt["P0"])
            idx.cell(r, 5).value = str(cnt["P1"])
            idx.cell(r, 6).value = str(cnt["P2+"])
            total.update(cnt)
        idx.cell(24, 3).value = str(total["n"])
        idx.cell(24, 4).value = str(total["P0"])
        idx.cell(24, 5).value = str(total["P1"])
        idx.cell(24, 6).value = str(total["P2+"])
        idx.cell(2, 1).value = (
            "Product: InternSafar (workspace sibling internship-portal). Re-synced 2026-09-28 against live routes/APIs: "
            "retired form/manual-request/Google-employer paths removed; employer email verify, Hybrid E documents, "
            "sticky approval, Suspend/Restore/Reject, Reset Ethics, Adjust Points added."
        )
        idx.cell(26, 2).value = f"Candidate: {CAND} / <IP_QA_CORE_PASSWORD> | Employer: {EMP} / <IP_QA_CORE_PASSWORD> | SuperAdmin: {SA} / <IP_QA_CORE_PASSWORD> (all on `/`)"
        stamp_index_for_results(wb)

    if "Coverage Matrix" in wb.sheetnames:
        cm = wb["Coverage Matrix"]
        rows = [
            ("Sign-in on `/` + captcha + 2FA (all roles; /superadmin/login redirects)", "Y", "Y", "Y", "Y"),
            ("Register: candidate Google Gmail / employer Domain or Free-email + email verify", "Y", "Y", "approve", "Y"),
            ("Employer login after email verify while pending (posting gated)", "—", "Y", "—", "—"),
            ("Candidate profile six tabs (phone required)", "Y", "—", "—", "—"),
            ("Browse tabs (Unapplied default) + List/Cards + filters Show/Hide + presets", "Y", "—", "—", "—"),
            ("Browse / save / apply (5 pts) + withdraw / re-apply", "Y", "—", "—", "—"),
            ("Employer profile five tabs + ethics Accept/Reject lock", "—", "Y", "Reset Ethics", "—"),
            ("Verification documents: one active per type (Hybrid E)", "—", "Y", "review", "—"),
            ("Final Approval gate (≥1 approved, 0 pending) + sticky approval", "—", "Y", "Y", "—"),
            ("Suspend / Restore / Reject / Delete employers", "—", "—", "Y", "—"),
            ("Postings (Publish last tab) + applicant pipeline + xlsx/zip exports", "—", "Y", "oversee", "—"),
            ("Offers / ratings / completions", "Y", "Y", "—", "—"),
            ("Messages + attachments + bulk archive", "Y", "Y", "oversight", "—"),
            ("Notifications UI", "Y", "Y", "no nav page", "—"),
            ("Points / referral /r/{code}", "Y", "Y", "Adjust Points", "landing"),
            ("Posting Share Rewards / Listing reports", "report", "claim", "Y", "—"),
            ("Feature ideas", "Y", "Y", "moderate", "—"),
            ("Account 2FA / sessions / email + phone change", "Y", "Y", "Y", "—"),
        ]
        cm.cell(1, 1).value = f"Coverage Matrix — synced {date.today().isoformat()} (live sibling)"
        for r in range(4, max(cm.max_row, 4 + len(rows)) + 1):
            for c in range(1, 6):
                cm.cell(r, c).value = None
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = cm.cell(4 + i, j + 1)
                cell.value = v
                style_cell(cell)

    stats["total_cases"] = total["n"]
    print(json.dumps(dict(stats), indent=2))
    print(f"RESIDUE ({len(residue)}):")
    for line in residue:
        print("  " + line)
    if not dry_run:
        wb.save(XLSX)
        print(f"saved {XLSX}")


if __name__ == "__main__":
    main(dry_run="--dry-run" in sys.argv)
