# -*- coding: utf-8 -*-
"""
Candidate profile step count (bug report 2026-10-09): only the four setup tabs are numbered
(1. Basics & Contact … 4. Work Readiness, header "Setup step X of 4"); Privacy & Photo and
Endorsements (Read-Only) have no number. Rewrites the case text in InternSafar-Test-Cases.xlsx
that still names "5. Privacy & Photo" / "6. Endorsements" or the old five-tab wizard.

Actual Result cells are left alone — they hold real run output and refresh on the next Excel apply.

  python scripts/sync-internsafar-xlsx-2026-10-09-profile-step-count.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-09-profile-step-count.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "06 Candidate Profile"

UPDATES = {
    "TC-IS-06-001": {
        "Comments / Notes": (
            "QASkills: EP complete vs incomplete; BVA phone blank vs invalid national number. Tabs today: "
            "1. Basics & Contact, 2. Academic, 3. Skills & Experience, 4. Work Readiness, then unnumbered "
            "Privacy & Photo and Endorsements (Read-Only). Setup steps 1–4 ('Setup step X of 4') on first visit. "
            "Legacy ID annotation if present stays.\n"
            "Legacy ID retained for history; Playwright apply uses TC-IS-* only."
        ),
    },
    "TC-IS-06-002": {
        "Description": (
            "1. Open /candidate/profile.\n2. Confirm tabs: 1. Basics & Contact, 2. Academic, 3. Skills & Experience, "
            "4. Work Readiness, then an 'Optional' divider, then Privacy & Photo and Endorsements (Read-Only) with no number. "
            "On 1. Basics & Contact the header reads 'Setup step 1 of 4'.\n"
            "3. On Privacy & Photo toggle show photo / searchable / hide phone until shortlist; set WhatsApp/Telegram "
            "opt-in.\n4. Save Privacy Settings and reload."
        ),
        "Suggestion / Expected Behaviour": (
            "All six tabs visible. Only the four setup tabs are numbered and show 'Setup step X of 4'; an 'Optional' "
            "divider separates them from Privacy & Photo and Endorsements, which have no number and no step header. "
            "Endorsements read-only. Privacy toggles and opt-ins persist via PUT /api/ip/candidate/profile."
        ),
    },
    "TC-IS-06-012": {
        "Description": (
            "1. Sign in as the test candidate and open /candidate/profile → Basics.\n"
            "2. LinkedIn Profile URL: type 'https://lnkd.in/abc123', then a company page "
            "'https://www.linkedin.com/company/acme'.\n"
            "3. GitHub / Portfolio URL: type 'priya'. Personal website: type 'javascript:alert(1)'.\n"
            "4. Open the Privacy & Photo tab → Instant Messaging Alerts: WhatsApp '12345', Telegram '@ab'.\n"
            "5. Click Save after each, then replace with valid values: 'linkedin.com/in/your-name', "
            "'github.com/your-name', '9876543210', 'your_name'."
        ),
    },
    "TC-IS-06-014": {
        "Description": (
            "1. Sign in as the test candidate and open /candidate/profile.\n"
            "2. Click the 'First Name' label — the First Name box gets focus. Type 'Priya2' (screen reader reads the error with the field).\n"
            "3. 4. Work Readiness → Preferred working hours: from 22:00, to 02:00 (night shift). Save.\n"
            "4. Ongoing commitment? = Other → Commitment note: try to type more than 200 characters.\n"
            "5. Open the Privacy & Photo tab → Telegram '@ab'. Go back to 1. Basics & Contact and click Save.\n"
            "6. Click 'Go to Privacy & Photo' in the save message."
        ),
        "Suggestion / Expected Behaviour": (
            "Step 2: label focuses the input; the name error is announced with it. "
            "Step 3: no error — night-shift hours (\"to\" earlier than \"from\") save as entered. "
            "Step 4: the box stops at 200 characters (API also returns 400 above 200). "
            "Step 5: 'Not saved — Telegram handle must be … (on the Privacy & Photo tab)' with a 'Go to Privacy & Photo' link. "
            "Step 6: the Privacy tab opens and the Telegram box is focused. "
            "PUT /api/ip/candidate/profile returns 400 only when an hours value is not a time (e.g. '9am') or the note is over 200 characters."
        ),
    },
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    changed = 0
    for case_id, fields in UPDATES.items():
        row = rows.get(case_id)
        if not row:
            raise SystemExit(f"{case_id} not found in '{SHEET}'")
        for key, text in fields.items():
            cell = ws.cell(row, cols[key])
            if cell.value == text:
                continue
            print(f"{'Updating' if apply else 'Would update'} {case_id} '{key}' (row {row})")
            cell.value = text
            changed += 1
    if not changed:
        print("Workbook text already current.")
        return
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name} ({changed} cell(s))")
    else:
        print(f"{changed} cell(s) would change — re-run with --apply")


if __name__ == "__main__":
    main("--apply" in sys.argv)
