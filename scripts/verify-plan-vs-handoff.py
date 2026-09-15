#!/usr/bin/env python3
"""One-shot plan vs zip/tar verification."""
import zipfile, tarfile, io
from pathlib import Path
from docx import Document

zpath = Path(r"C:\Users\place\Work\UIUX Migration\internship-portal-aws-handoff.zip")
app = Path(r"C:\Users\place\Work\UIUX Migration\internship-portal")
results = []


def ok(name, cond, detail=""):
    results.append((bool(cond), name, detail))


with zipfile.ZipFile(zpath) as z:
    names = set(z.namelist())
    tarname = [n for n in names if n.endswith(".tar.gz")][0]
    ok("ZIP has app tar.gz", True, tarname)
    for f in [
        "README.txt",
        "PATH-B-NO-DB-MIGRATE.txt",
        "APP-TAR-FOLDER-STRUCTURE.txt",
        "AWS-DEPLOY.md",
        "ISSUES-FIXED.txt",
        "SETUP-NOTES.txt",
        "DB-REFERENCE.md",
        "DB-SCRIPTS-REFERENCE.txt",
        "runner/assert-db-migrate-allowed.js",
        "runner/db_exec_sql_file.js",
        "runner/db_migrate_sql_only_ip.mjs",
        "runner/MIGRATION_MANIFEST.txt",
    ]:
        ok(f"ZIP has {f}", f in names)

    mig = [n for n in names if n.startswith("migrations/") and n.endswith(".sql")]
    ok("ZIP migrations count 40", len(mig) == 40, str(len(mig)))
    ok("ZIP has migration 039", any("039_ip_feature" in n for n in names))
    ok("ZIP NO db_migrate_all", not any("db_migrate_all" in n for n in names))

    readme = z.read("README.txt").decode()
    ok("README IP_ALLOW gate", "IP_ALLOW_DB_MIGRATE=1" in readme)
    ok("README fail-closed OK/FAIL", "=== FAIL" in readme and "=== OK" in readme)
    pathb = z.read("PATH-B-NO-DB-MIGRATE.txt").decode()
    ok("PATH-B why + code gate", "assert-db-migrate-allowed" in pathb and "Why" in pathb)

    data = z.read(tarname)
    with tarfile.open(fileobj=io.BytesIO(data), mode="r:gz") as tar:
        tnames = set(tar.getnames())
        ok("TAR forward slashes", all("\\" not in n for n in tnames))

        def has_end(suffix):
            return any(n.endswith(suffix) for n in tnames)

        for s in [
            "APP-FOLDER-STRUCTURE.txt",
            "AGENTS.md",
            "deploy-fresh-aws-db.mjs",
            "assert-db-migrate-allowed.js",
            "db_exec_sql_file.js",
            "ipSeedCoreBaseline.js",
            "fill-core-coverage.mjs",
            ".env.example",
            "package.json",
        ]:
            ok(f"TAR has {s}", has_end(s))

        ok("TAR migration 039", any("039_ip_feature" in n for n in tnames))
        ok("TAR auth.js", any(n.endswith("src/lib/auth.js") for n in tnames))
        ok("TAR PortalShell.jsx", any("PortalShell.jsx" in n for n in tnames))
        ok("TAR IpSignInLanding.jsx", any("IpSignInLanding.jsx" in n for n in tnames))
        ok("TAR NO migrate-and-seed.mjs", not any(n.endswith("migrate-and-seed.mjs") for n in tnames))
        ok("TAR NO .env", not any(n == ".env" or n.endswith("/.env") for n in tnames))

        def tread(endswith):
            m = [n for n in tnames if n.endswith(endswith)][0]
            return tar.extractfile(m).read().decode("utf-8", "replace")

        seed = tread("ipSeedCoreBaseline.js")
        text_casts = seed.count("::text[]")
        ok("Seed ::text[] casts", text_casts >= 2, f"count={text_casts}")
        ok("Seed referral_code handled", "referral_code" in seed)

        fill = tread("fill-core-coverage.mjs")
        ok(
            "fill-core column detect",
            "problem" in fill
            and ("information_schema" in fill.lower() or "hasColumn" in fill or "column" in fill.lower()),
        )

        dep = tread("deploy-fresh-aws-db.mjs")
        ok("Orchestrator gate", "assert-db-migrate-allowed" in dep)
        ok("Orchestrator 035 split", "035_ip_seed_candidate_academics" in dep)
        ok("Orchestrator IP_Reset", "IP_Reset_Core_Sample" in dep)

        auth = tread("src/lib/auth.js")
        ok("Auth trustHost", "trustHost" in auth)
        ok("Auth signOut event", "signOut" in auth)

        agents = tread("AGENTS.md")
        ok("AGENTS IP_ALLOW + Path B", "IP_ALLOW_DB_MIGRATE" in agents and "Path B" in agents)

        pkg = tread("package.json")
        ok("package.json deploy script", "deploy:fresh-aws-db" in pkg)

        gate = tread("assert-db-migrate-allowed.js")
        ok("Gate BLOCKED banner", "=== BLOCKED" in gate)

        db_exec = tread("db_exec_sql_file.js")
        ok("db_exec calls gate", "assert-db-migrate-allowed" in db_exec)
        ok("db_exec FAIL banner", "=== FAIL" in db_exec)

manifest = (app / "scripts" / "MIGRATION_MANIFEST.txt").read_text(encoding="utf-8")
ok("Live manifest has 039", "039_ip_feature_idea_detail_columns.sql" in manifest)
ok("Live captchaBypass.js", (app / "src" / "lib" / "captchaBypass.js").exists())

doc = Document(
    r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
)
dtext = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for c in row.cells:
            dtext += "\n" + c.text
ok("Docx Part S", "Part S" in dtext and "assert-db-migrate-allowed" in dtext)
ok("Docx Part M", "Part M" in dtext)
ok("Docx IP_ALLOW", "IP_ALLOW_DB_MIGRATE=1" in dtext)
ok("Docx no db:migrate:all leftover", "db:migrate:all" not in dtext)

print("=== PLAN vs ZIP/TAR ===")
fail = 0
for good, name, detail in results:
    mark = "PASS" if good else "FAIL"
    if not good:
        fail += 1
    print(f"{mark}: {name}" + (f" ({detail})" if detail else ""))
print(f"\nTOTAL {len(results)}  FAIL {fail}")
raise SystemExit(1 if fail else 0)
