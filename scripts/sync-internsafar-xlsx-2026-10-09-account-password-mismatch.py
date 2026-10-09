# -*- coding: utf-8 -*-
"""
Add TC-IS-05-007 (Account → Change Permanent Password: an empty or different Confirm New Password
shows "Passwords do not match" in the form) to InternSafar-Test-Cases.xlsx.

Bug report 2026-10-09 #3: with Confirm New Password left blank, Update Password showed only the
browser's own "Please fill out this field." bubble and no mismatch message. The form now runs its
own checks (noValidate). Automated by Playwright IS-086 on the test candidate; no password is changed.
Re-running with --apply updates the row if it already exists.

  python scripts/sync-internsafar-xlsx-2026-10-09-account-password-mismatch.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-09-account-password-mismatch.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "05 Account & Security"
TEMPLATE_ID = "TC-IS-05-006"

NEW_CASE = {
    "ID": "TC-IS-05-007",
    "Module / Section": "05 Account & Security",
    "Type": "Negative",
    "Issue Summary": "Change password shows 'Passwords do not match' when Confirm New Password is empty or different",
    "Description": (
        "1. Sign in as the test candidate and open /account → Security & Password.\n"
        "2. Fill Current Password and New Password (e.g. 'Qa!Mismatch123'); leave Confirm New Password empty; click Update Password.\n"
        "3. Type a different value in Confirm New Password (e.g. 'Qa!Mismatch124'); click Update Password.\n"
        "4. Type the same value as New Password in Confirm New Password.\n"
        "5. Click Clear Form."
    ),
    "Suggestion / Expected Behaviour": (
        "Step 2: '✗ Passwords do not match' under Confirm New Password and 'New passwords do not match. Re-enter your new "
        "password in Confirm New Password.' above the buttons; the confirm box is focused and marked invalid "
        "(no browser-only 'Please fill out this field.' bubble). "
        "Step 3: the bottom message clears while typing, then 'New passwords do not match.' after Update Password. "
        "Step 4: '✓ Passwords match' and no error. Step 5: both messages gone. "
        "No request reaches /api/ip/auth/change-password in steps 2–5."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P2",
    "Test Status": "Not Run",
    "Comments / Notes": "Bug report 2026-10-09 #3 (password mismatch not shown). Playwright IS-086 on the test candidate; password never changed.",
    "Phase": "2026-10-09 bug report",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Change password validation",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        for key in ("Issue Summary", "Description", "Suggestion / Expected Behaviour", "Feature"):
            if key in cols:
                ws.cell(row, cols[key]).value = NEW_CASE[key]
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
