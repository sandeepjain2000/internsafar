#!/usr/bin/env python3
"""Append Part R — Path B/C DB commands + CAPTCHA bypass (current only)."""
from docx import Document
from docx.enum.text import WD_BREAK

DOC_PATH = r"C:\Users\place\Work\UIUX Migration\InternSafar_AWS_Deployment_Runbook_CRISP_COMPLETE_FINAL (1) (1).docx"
MARKER = "Part R — Final DB script decision + CAPTCHA bypass"


def main():
    doc = Document(DOC_PATH)
    if MARKER in "\n".join(p.text for p in doc.paragraphs):
        print("Part R already present — skip")
        return

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(MARKER, level=1)

    doc.add_heading("Database script routing", level=2)
    doc.add_paragraph("Use only the commands below:")
    for line in [
        "Path B (app update, RDS unchanged): do NOT run any database migrate or seed.",
        "Path C (fresh/empty RDS): npm run deploy:fresh-aws-db only.",
        "SQL only when demo users already exist: npm run db:migrate:sql-only.",
    ]:
        doc.add_paragraph("• " + line)

    doc.add_heading("Fail-closed migrations", level=2)
    doc.add_paragraph(
        "Runners print === FAIL === and exit 1 on error; success requires === OK === and exit 0. "
        "Do not claim success if a SQL file failed."
    )

    doc.add_heading("CAPTCHA bypass — why it is true, how to enable real captcha", level=2)
    doc.add_paragraph(
        "This build ships CAPTCHA_BYPASS_FOR_TESTING = true in src/lib/captchaBypass.js."
    )
    doc.add_paragraph(
        "Why: on AWS, missing NEXTAUTH_SECRET made createLoginCaptcha throw → "
        "/api/auth/captcha returned HTTP 500 → UI “Verification unavailable”. Bypass was a "
        "temporary way to keep login usable."
    )
    doc.add_paragraph("Still set these on EC2 (bypass does not replace them):")
    for line in [
        "NEXTAUTH_URL=https://internsafar.com",
        "NEXTAUTH_SECRET=<openssl rand -base64 32 or value provided separately>",
        "Then: npm run build && pm2 restart internsafar --update-env",
    ]:
        doc.add_paragraph("• " + line)
    doc.add_paragraph("To turn real numbered CAPTCHA back on later:")
    for line in [
        "curl -i http://127.0.0.1:3000/api/auth/captcha must return HTTP 200",
        "Set CAPTCHA_BYPASS_FOR_TESTING = false",
        "Path B redeploy (build + PM2)",
    ]:
        doc.add_paragraph("• " + line)

    doc.save(DOC_PATH)
    print(f"Updated: {DOC_PATH}")


if __name__ == "__main__":
    main()
