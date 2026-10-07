# -*- coding: utf-8 -*-
"""
Sync InternSafar-Test-Cases.xlsx to the 2026-10-07 test-account rule.

- Core demo accounts are never used for testing: Preconditions / steps / notes / Index
  now name the disposable test accounts from scripts/lib/ipTestAccountsConfig.js
  (re-created by `npm run qa:ensure-test-accounts`). SuperAdmin stays shared.
- Wrong-role rows expect the "Wrong account for this workspace" block page
  (URL stays, no other-role nav) instead of a redirect / bounce.
- Actual Result history is left as recorded. TC-IS-18-034 (bootstrap keeps the core
  accounts) still names the cores on purpose.

  python scripts/sync-internsafar-xlsx-2026-10-07-test-accounts.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-07-test-accounts.py --apply   # writes workbook
"""
from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"

TEST_CAND = "lawsonlclintern+qa1@gmail.com"
TEST_EMP = "placementhubsupport+qa1@gmail.com"
TEST_EMP_PENDING = "placementhubsupport+qa2@gmail.com"

CORE_RE = re.compile(r"lawsonlclintern(\+|%2B)1(@|%40)|placementhubsupport(@|%40)gmail|\bcore (candidate|employer)\b|Nova Labs", re.I)
SKIP_COLUMNS = {"Actual Result"}
KEEP_CORE_IDS = {"TC-IS-18-034"}

EMAIL_SWAPS = [
    (re.compile(r"lawsonlclintern(\+|%2B)1(@|%40)gmail\.com", re.I),
     lambda m: f"lawsonlclintern{m.group(1)}qa1{m.group(2)}gmail.com"),
    (re.compile(r"placementhubsupport(\+|%2B)3(@|%40)gmail\.com", re.I),
     lambda m: f"placementhubsupport{m.group(1)}qa2{m.group(2)}gmail.com"),
    (re.compile(r"placementhubsupport(@|%40)gmail\.com", re.I),
     lambda m: f"placementhubsupport{'+' if m.group(1) == '@' else '%2B'}qa1{m.group(1)}gmail.com"),
]

TEXT_SWAPS = [
    (re.compile(r"\bcore candidate\b", re.I), "test candidate"),
    (re.compile(r"\bcore employer\b", re.I), "test employer"),
    (re.compile(r"\bNova Labs\b"), "Kestrel Analytics"),
]

TEST_ACCOUNT_NOTE = (
    " Test accounts only (never the core demo accounts); run `npm run qa:ensure-test-accounts` "
    "first if they are missing."
)

ROW_REWRITES = {
    "TC-IS-04-002": {
        "Issue Summary": "Wrong-role workspace shows the 'Wrong account' block, not the other role's UI",
        "Suggestion / Expected Behaviour": (
            "Signed in as one role and opening another role's workspace (e.g. candidate on /employer) shows "
            "'Wrong account for this workspace' with the signed-in email and a Sign out button. The URL may stay "
            "on the other role's path, but none of that role's navigation or data renders (403-style block, "
            "no redirect required)."
        ),
    },
    "TC-IS-14-001": {
        "Suggestion / Expected Behaviour": (
            "Stats load for SuperAdmin only; counts match the live queues (no Form Registrations / Manual "
            "Requests). Candidate opening /superadmin gets the 'Wrong account for this workspace' block "
            "(no SuperAdmin nav or stats)."
        ),
    },
    "TC-IS-18-048": {
        "Description": (
            "1. Sign in as candidate / employer / SuperAdmin (test accounts); open role home; reload.\n"
            "2. As the test candidate, open /employer and reload; confirm the 'Wrong account for this "
            "workspace' block shows and no employer navigation renders."
        ),
    },
}


def header_map(ws):
    for row in ws.iter_rows(min_row=1, max_row=3):
        names = [str(c.value).strip() if c.value is not None else "" for c in row]
        if "ID" in names or "TC ID" in names:
            return row[0].row, {n: i for i, n in enumerate(names) if n}
    return None, {}


def swap_text(text, keep_core):
    out = text
    if not keep_core:
        for pat, repl in EMAIL_SWAPS:
            out = pat.sub(repl, out)
        for pat, repl in TEXT_SWAPS:
            out = pat.sub(repl, out)
    return out


def main(apply=False):
    wb = load_workbook(XLSX)
    changed = Counter()
    rewrites_done = []
    emails_after = Counter()
    email_re = re.compile(r"[\w.+%-]+(?:@|%40)[\w.-]+\.\w+")
    core_left = []

    for ws in wb.worksheets:
        hdr_row, cols = header_map(ws)
        id_idx = cols.get("ID", cols.get("TC ID"))
        start = (hdr_row or 0) + 1
        for row in ws.iter_rows(min_row=start):
            tc_id = str(row[id_idx].value).strip() if id_idx is not None and row[id_idx].value else ""
            keep_core = tc_id in KEEP_CORE_IDS
            for name, idx in cols.items():
                if name in SKIP_COLUMNS or idx >= len(row):
                    continue
                cell = row[idx]
                if not isinstance(cell.value, str):
                    continue
                new = swap_text(cell.value, keep_core)
                if name == "Preconditions" and new != cell.value and TEST_CAND in new \
                        and "qa:ensure-test-accounts" not in new:
                    new = new.rstrip() + TEST_ACCOUNT_NOTE
                if new != cell.value:
                    changed[(ws.title if ws.title == "Index" else "cases", name)] += 1
                    cell.value = new
            for name, text in ROW_REWRITES.get(tc_id, {}).items():
                if name in cols and row[cols[name]].value != text:
                    row[cols[name]].value = text
                    rewrites_done.append((tc_id, name))
        if hdr_row is None:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        new = swap_text(cell.value, False)
                        if new != cell.value:
                            changed[(ws.title, "free text")] += 1
                            cell.value = new
        names_by_idx = {i: n for n, i in cols.items()}
        for row in ws.iter_rows():
            row_id = str(row[id_idx].value) if id_idx is not None and row[id_idx].value else ""
            for i, cell in enumerate(row):
                if isinstance(cell.value, str):
                    for m in email_re.finditer(cell.value):
                        emails_after[m.group(0)] += 1
                    if CORE_RE.search(cell.value):
                        core_left.append((ws.title, row_id, names_by_idx.get(i, f"col{i}")))

    print("Cells changed (sheet group, column):")
    for k, n in sorted(changed.items()):
        print(f"  {n:4d}  {k}")
    print("Row rewrites:", rewrites_done or "none")
    print("Emails remaining after sync (all columns incl. Actual Result history):")
    for e, n in emails_after.most_common():
        print(f"  {n:4d}  {e}")
    print("Core-account mentions left (expected: Actual Result history, TC-IS-18-034):")
    for loc in core_left:
        print("  ", loc)
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX}")
    else:
        print("Preview only. Re-run with --apply to write.")


if __name__ == "__main__":
    main(apply="--apply" in sys.argv)
