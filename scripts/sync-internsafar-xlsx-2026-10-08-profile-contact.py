# -*- coding: utf-8 -*-
"""
Add TC-IS-06-012 (candidate profile links, WhatsApp and Telegram are validated) to
InternSafar-Test-Cases.xlsx — follow-up to the external QA report recommendation
"consistent validation rules for all Candidate personal-information fields".

Rule (src/lib/ipProfileContact.js): LinkedIn must be a linkedin.com/in/… profile link;
GitHub / Portfolio, Portfolio and Personal website must be http(s) web links (bare domains get
https://); 255 characters max; WhatsApp must be a valid number for the mobile country code;
Telegram 5–32 letters/numbers/underscores starting with a letter (stored with @). Only values
changed since the last save are checked. Inline on /candidate/profile and by
PUT /api/ip/candidate/profile (400). Automated by Playwright regression IS-080.

  python scripts/sync-internsafar-xlsx-2026-10-08-profile-contact.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-08-profile-contact.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "06 Candidate Profile"
TEMPLATE_ID = "TC-IS-06-011"

NEW_CASE = {
    "ID": "TC-IS-06-012",
    "Module / Section": "06 Candidate Profile",
    "Type": "Negative",
    "Issue Summary": "LinkedIn, GitHub / Portfolio, website, WhatsApp and Telegram are validated",
    "Description": (
        "1. Sign in as the test candidate and open /candidate/profile → Basics.\n"
        "2. LinkedIn Profile URL: type 'https://lnkd.in/abc123', then a company page "
        "'https://www.linkedin.com/company/acme'.\n"
        "3. GitHub / Portfolio URL: type 'priya'. Personal website: type 'javascript:alert(1)'.\n"
        "4. 5. Privacy & Photo → Instant Messaging Alerts: WhatsApp '12345', Telegram '@ab'.\n"
        "5. Click Save after each, then replace with valid values: 'linkedin.com/in/your-name', "
        "'github.com/your-name', '9876543210', 'your_name'."
    ),
    "Suggestion / Expected Behaviour": (
        "Each invalid value shows an inline message under its field and Save shows 'Not saved — …'; "
        "nothing is stored. LinkedIn: 'Enter your LinkedIn profile link, like https://www.linkedin.com/in/your-name.' "
        "Links: '<Field> must be a web link, like https://yourname.com.' WhatsApp: 'WhatsApp number isn't valid for +91…'. "
        "Telegram: 'Telegram handle must be 5–32 letters, numbers or underscores, starting with a letter…'. "
        "PUT /api/ip/candidate/profile returns 400 with the same message. Valid values save; bare links are stored "
        "with https:// and Telegram with a leading @. A previously saved value that was never edited does not block Save."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P2",
    "Test Status": "Not Run",
    "Comments / Notes": "From external QA report recommendations (2026-10-08). Playwright IS-080.",
    "Phase": "2026-10-08 profile contact validation",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Profile link / handle validation",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        print(f"{NEW_CASE['ID']} already present (row {rows[NEW_CASE['ID']]}) — nothing to add.")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
