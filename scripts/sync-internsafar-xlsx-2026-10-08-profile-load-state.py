# -*- coding: utf-8 -*-
"""
Add TC-IS-06-013 (candidate profile shows a loading state, and an error state with Try again
when the profile request fails) to InternSafar-Test-Cases.xlsx.

/candidate/profile: spinner "Loading your profile…" while GET /api/ip/candidate/profile runs;
on failure (after quiet retries) a "We couldn't load your profile" panel with Try again; on 401
a "Your session has expired" panel with Sign in again. Automated by Playwright regression IS-081.

  python scripts/sync-internsafar-xlsx-2026-10-08-profile-load-state.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-08-profile-load-state.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "06 Candidate Profile"
TEMPLATE_ID = "TC-IS-06-012"

NEW_CASE = {
    "ID": "TC-IS-06-013",
    "Module / Section": "06 Candidate Profile",
    "Type": "Negative",
    "Issue Summary": "Profile page shows a loading state, and an error with Try again when loading fails",
    "Description": (
        "1. Sign in as the test candidate.\n"
        "2. Open /candidate/profile on a slow connection (DevTools → Network → Slow 3G).\n"
        "3. Switch DevTools Network to Offline and reload /candidate/profile.\n"
        "4. Switch back to Online and click Try again."
    ),
    "Suggestion / Expected Behaviour": (
        "While loading: a spinner with 'Loading your profile…' (no blank page, no form). "
        "When the request fails: a panel 'We couldn't load your profile' with a plain reason "
        "(e.g. 'We could not reach the server. Check your internet connection.') and a Try again button; "
        "no half-empty form is shown. Try again loads the profile form. If the session has expired (401), "
        "the panel says 'Your session has expired' with Sign in again."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P2",
    "Test Status": "Not Run",
    "Comments / Notes": "From external QA report follow-up (2026-10-08). Playwright IS-081.",
    "Phase": "2026-10-08 profile load state",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Profile loading / error state",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        print(f"{NEW_CASE['ID']} already present (row {rows[NEW_CASE['ID']]}) — nothing to add.")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
