# -*- coding: utf-8 -*-
"""
Whole-workbook QA coverage: accounts for every TC-IS case in InternSafar-Test-Cases.xlsx.

A test run is only "all tests" when this report says so. It splits the workbook into
  - automated (checklist runner, Playwright mapping, or SCRIPT_RUNNERS) -> verified on/after --since, or NOT run
  - manual only (no test code) -> needs a person
and the two always add up to the workbook total.

  npm run qa:coverage                      # since = today (IST)
  python scripts/report-internsafar-qa-coverage.py --since 2026-10-07 --strict
  npm run qa:coverage -- --write-map --sync-automation   # also regenerate the case map + Automation column

--strict exits 1 when any automated case was not verified on/after --since.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
RUNNER_RESULTS = ROOT / "test-cases" / "qa-results.json"
PLAYWRIGHT_MAP = ROOT / "scripts" / "apply-playwright-regression-xlsx.mjs"
CASE_MAP = ROOT / "docs" / "ai-context" / "domains" / "testing-case-map.md"
TC_RE = re.compile(r"TC-IS-\d\d-\d{3}")

# Automated by scripts outside the checklist runner / Playwright. They record into qa-results.json
# (scripts/lib/recordQaResults.mjs), so their ids must not be counted as checklist-runner rows.
SCRIPT_RUNNERS = {
    "TC-IS-03-013": "scripts/manual/run-tc-is-03-013-duplicate-employer.mjs",
    "TC-IS-03-015": "scripts/manual/run-tc-is-03-015-self-referral.mjs",
    "TC-IS-03-022": "scripts/manual/run-tc-is-03-022-register-reject.mjs",
    "TC-IS-06-007": "scripts/manual/run-tc-is-06-007-email-change.mjs",
    "TC-IS-03-026": "scripts/qa-employer-reg-verify-approve-login.mjs",
    "TC-IS-03-030": "scripts/qa-employer-reg-verify-approve-login.mjs",
    "TC-IS-14-024": "scripts/qa-employer-reg-verify-approve-login.mjs",
    "TC-IS-18-051": "scripts/qa-employer-reg-verify-approve-login.mjs",
    "TC-IS-02-028": "scripts/qa-register-approve-post-apply-smoke.mjs",
    "TC-IS-02-029": "scripts/qa-register-approve-post-apply-smoke.mjs",
    "TC-IS-03-025": "scripts/qa-register-approve-post-apply-smoke.mjs",
    **{
        tc: "scripts/qa-test-account-cases.mjs"
        for tc in (
            "TC-IS-03-027", "TC-IS-09-019", "TC-IS-09-020", "TC-IS-09-021", "TC-IS-12-011",
            "TC-IS-12-012", "TC-IS-13-005", "TC-IS-14-029", "TC-IS-14-032", "TC-IS-17-007",
            "TC-IS-18-053",
            "TC-IS-05-001", "TC-IS-05-002", "TC-IS-06-002", "TC-IS-07-003", "TC-IS-09-006",
            "TC-IS-11-001", "TC-IS-11-005", "TC-IS-12-001", "TC-IS-12-002", "TC-IS-12-004",
            "TC-IS-13-001", "TC-IS-14-010", "TC-IS-14-011", "TC-IS-14-014", "TC-IS-14-015",
            "TC-IS-14-016", "TC-IS-14-017", "TC-IS-15-003", "TC-IS-18-012", "TC-IS-18-015",
            "TC-IS-18-016", "TC-IS-18-017", "TC-IS-18-020",
        )
    },
    **{
        tc: "scripts/qa-temp-employer-cases.mjs"
        for tc in (
            "TC-IS-14-005", "TC-IS-14-025", "TC-IS-14-026", "TC-IS-14-027", "TC-IS-14-028",
            "TC-IS-14-033", "TC-IS-17-006", "TC-IS-18-019",
        )
    },
}
# Scripted, but only reach Pass/Fail with extra setup; otherwise they report Blocked, which the
# workbook writer does not store for these ids. Listed apart so --strict does not fail on them.
NEEDS_SETUP = {
    "TC-IS-03-008": "live Google OAuth consent by a person (checklist runner reports Blocked)",
    "TC-IS-03-015": "local server with IP_ALLOW_UNVERIFIED_GOOGLE_REGISTER=1",
    "TC-IS-06-007": "a person pastes the emailed code into IP_QA_EMAIL_CHANGE_CODE",
}


def today_ist() -> str:
    return (datetime.now(timezone.utc) + timedelta(hours=5, minutes=30)).strftime("%Y-%m-%d")


def read_workbook():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    cases = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = None
        for row in rows:
            cells = [str(c).strip() if c is not None else "" for c in (row or ())]
            if "ID" in cells:
                header = cells
                break
        if not header:
            continue
        col = {
            name: header.index(name)
            for name in ("ID", "Test Status", "Date Verified", "Automation", "Legacy ID")
            if name in header
        }
        for row in rows:
            if not row:
                continue
            tc = row[col["ID"]] if col["ID"] < len(row) else None
            if not (isinstance(tc, str) and TC_RE.fullmatch(tc.strip())):
                continue

            def cell(name):
                i = col.get(name)
                return row[i] if i is not None and i < len(row) else None

            cases[tc.strip()] = {
                "sheet": ws.title,
                "status": str(cell("Test Status") or "Not Run").strip(),
                "date": str(cell("Date Verified") or "")[:10],
                "automation": str(cell("Automation") or "").strip(),
                "legacy": str(cell("Legacy ID") or "").strip(),
            }
    return cases


def runner_result_keys() -> set[str]:
    if not RUNNER_RESULTS.exists():
        return set()
    data = json.loads(RUNNER_RESULTS.read_text(encoding="utf-8"))
    return set((data.get("cases") or {}).keys()) | set((data.get("byTcId") or {}).keys())


def runner_ids(cases) -> set[str]:
    """Rows the checklist runner covers: by TC-IS id, or by its older checklist id in the Legacy ID column."""
    keys = runner_result_keys()
    return {
        tc
        for tc, row in cases.items()
        if tc not in SCRIPT_RUNNERS and (tc in keys or (row["legacy"] and row["legacy"] in keys))
    }


def playwright_ids() -> set[str]:
    return set(TC_RE.findall(PLAYWRIGHT_MAP.read_text(encoding="utf-8")))


def write_case_map(cases, runner, pw, scripted) -> None:
    result_keys = runner_result_keys()
    lines = [
        "# Testing — every workbook case and the runner that covers it (generated)",
        "",
        "Regenerate: `npm run qa:coverage -- --write-map` (`scripts/report-internsafar-qa-coverage.py`).",
        "Status / Date are what the workbook holds now. Runners and commands: `domains/testing.md` → Runner map.",
        "Every runner listed here writes its result to the workbook; \"needs setup\" rows only reach Pass/Fail with that setup.",
        "",
        "| ID | Sheet | Covered by | Status | Date |",
        "|---|---|---|---|---|",
    ]
    for tc in sorted(cases):
        row = cases[tc]
        by = []
        if tc in runner:
            by.append("checklist runner" + (f" (as {row['legacy']})" if tc not in result_keys and row["legacy"] else ""))
        if tc in pw:
            by.append("Playwright")
        if tc in scripted:
            by.append(f"`{SCRIPT_RUNNERS[tc]}`" + (f" (needs setup: {NEEDS_SETUP[tc]})" if tc in NEEDS_SETUP else ""))
        lines.append(
            f"| {tc} | {row['sheet']} | {'; '.join(by) or 'manual (no test code)'} | {row['status']} | {row['date'] or '—'} |"
        )
    CASE_MAP.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {CASE_MAP.relative_to(ROOT)}")


def sync_automation_column(automated) -> None:
    """Workbook Automation column = Automated when test code covers the id, else Manual (Obsolete kept)."""
    wb = openpyxl.load_workbook(WORKBOOK)
    changed = 0
    for ws in wb.worksheets:
        header_row = None
        for r in range(1, 12):
            vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
            if "ID" in vals and "Automation" in vals:
                header_row, id_col, auto_col = r, vals.index("ID") + 1, vals.index("Automation") + 1
                break
        if not header_row:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            tc = str(ws.cell(r, id_col).value or "").strip()
            cur = str(ws.cell(r, auto_col).value or "").strip()
            if not TC_RE.fullmatch(tc) or cur == "Obsolete":
                continue
            want = "Automated" if tc in automated else "Manual"
            if cur != want:
                ws.cell(r, auto_col).value = want
                changed += 1
    wb.save(WORKBOOK)
    print(f"Automation column synced ({changed} row(s) changed)")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--since", default=today_ist(), help="YYYY-MM-DD; results older than this count as not run")
    ap.add_argument("--strict", action="store_true", help="exit 1 if any automated case was not verified since --since")
    ap.add_argument("--write-map", action="store_true", help=f"write every case -> runner to {CASE_MAP.relative_to(ROOT)}")
    ap.add_argument("--sync-automation", action="store_true", help="set the workbook Automation column from test coverage")
    args = ap.parse_args()

    cases = read_workbook()
    runner = runner_ids(cases)
    pw = playwright_ids() & cases.keys()
    scripted = set(SCRIPT_RUNNERS) & cases.keys()
    needs_setup = set(NEEDS_SETUP) & cases.keys()
    automated = runner | pw | scripted
    manual = cases.keys() - automated
    if args.write_map:
        write_case_map(cases, runner, pw, scripted)
    if args.sync_automation:
        sync_automation_column(automated)

    fresh = {tc for tc in automated if cases[tc]["date"] >= args.since}
    stale = sorted(automated - fresh - needs_setup)
    setup_not_run = sorted(needs_setup - fresh)
    fresh_status = Counter(cases[tc]["status"] for tc in fresh)

    manual_by_sheet = defaultdict(list)
    for tc in sorted(manual):
        manual_by_sheet[cases[tc]["sheet"]].append(tc)
    manual_recent = sum(1 for tc in manual if cases[tc]["date"] >= args.since)

    total = len(cases)
    print(f"InternSafar workbook coverage (results on/after {args.since})")
    print(f"  Workbook TC-IS cases: {total}")
    print(
        f"  Automated: {len(automated)}  (checklist runner {len(runner)}, Playwright {len(pw)}, both {len(runner & pw)}, "
        f"other scripts {len(scripted)})"
    )
    print(
        f"    verified since {args.since}: {len(fresh)}  -> "
        + (", ".join(f"{k} {v}" for k, v in sorted(fresh_status.items())) or "none")
    )
    print(f"    NOT run since {args.since}: {len(stale)}")
    if stale:
        print("      " + ", ".join(stale))
    if setup_not_run:
        print(f"    Scripted but need setup to give Pass/Fail (not counted by --strict): {len(setup_not_run)}")
        for tc in setup_not_run:
            print(f"      {tc}: {NEEDS_SETUP[tc]}")
    print(f"  Manual only (no test code, needs a person): {len(manual)}  (result recorded since {args.since}: {manual_recent})")
    for sheet, ids in manual_by_sheet.items():
        print(f"    {sheet}: {len(ids)}")
    print(f"  Accounted: {len(automated) + len(manual)} of {total}")

    if len(fresh) == len(automated) and not manual_recent and manual:
        print(f"\nAll automated cases ran. {len(manual)} manual cases were NOT tested in this pass.")
    if args.strict and stale:
        print(f"\nFAIL: {len(stale)} automated cases were not run since {args.since}.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
