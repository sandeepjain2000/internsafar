# -*- coding: utf-8 -*-
"""
Add TC-IS-02-030 (after Sign out, the browser Back button returns to the sign-in page instead of a
stuck "Signing out…" screen) to InternSafar-Test-Cases.xlsx.

Bug report 2026-10-09 #4: after signing out, Back showed the restored portal page frozen on
"Signing out…". Real browsers restore that page from the back/forward cache; PortalShell now checks
the session on restore and sends the user to sign-in. Automated by Playwright IS-087 on the test
candidate with the back/forward cache switched on (Playwright turns it off by default).
Re-running with --apply updates the row if it already exists.

  python scripts/sync-internsafar-xlsx-2026-10-09-signout-back.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-09-signout-back.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "02 Auth & Access"
TEMPLATE_ID = "TC-IS-02-029"

NEW_CASE = {
    "ID": "TC-IS-02-030",
    "Module / Section": "02 Auth & Access",
    "Type": "Negative",
    "Issue Summary": "Browser Back after Sign out returns to sign-in, not a stuck 'Signing out…' page",
    "Description": (
        "1. Sign in as the test candidate and open /candidate/offers.\n"
        "2. Open /help, then press the browser Back button.\n"
        "3. Click Sign out and wait for the sign-in page.\n"
        "4. Press the browser Back button.\n"
        "5. Type /candidate/messages in the address bar."
    ),
    "Suggestion / Expected Behaviour": (
        "Step 2: back on /candidate/offers, still signed in. "
        "Step 3: sign-in page at /. "
        "Step 4: sign-in page at /?next=%2Fcandidate%2Foffers; 'Signing out…' is never left on screen. "
        "Step 5: sign-in page at /?next=%2Fcandidate%2Fmessages."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P1",
    "Test Status": "Not Run",
    "Comments / Notes": (
        "Bug report 2026-10-09 #4 (Back after sign-out stuck on 'Signing out…'). Playwright IS-087 on the test "
        "candidate, with the browser back/forward cache switched on (Playwright disables it by default)."
    ),
    "Phase": "2026-10-09 bug report",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "Sign out / session guard",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        for key in ("Issue Summary", "Description", "Suggestion / Expected Behaviour", "Feature"):
            if key in cols:
                ws.cell(row, cols[key]).value = NEW_CASE[key]
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
