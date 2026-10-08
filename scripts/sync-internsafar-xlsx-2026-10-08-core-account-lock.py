# -*- coding: utf-8 -*-
"""
Add TC-IS-05-006 (core / shared demo account keeps its password) to InternSafar-Test-Cases.xlsx.

Rule: ip_users.is_core_account (src/lib/ipCoreAccount.js). For flagged accounts, POST
/api/ip/auth/change-password and POST /api/ip/auth/password-reset/confirm run every check and return
the normal success reply, but the stored password and other sessions stay unchanged. Passwords for
core accounts are set only with scripts/ip-core-account.mjs. Automated by Playwright IS-083 on a
disposable test account (never the real core accounts). Re-running with --apply updates the row.

  python scripts/sync-internsafar-xlsx-2026-10-08-core-account-lock.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-08-core-account-lock.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "05 Account & Security"
TEMPLATE_ID = "TC-IS-05-005"

NEW_CASE = {
    "ID": "TC-IS-05-006",
    "Module / Section": "05 Account & Security",
    "Type": "Negative",
    "Issue Summary": "Core (shared demo) account keeps its password after a UI change or reset link; normal accounts still change",
    "Description": (
        "1. Backend: flag the second test candidate as core — npm run ip:core-account -- --mark=<test candidate email>.\n"
        "2. Sign in as that test candidate → /account → Change password with a wrong current password.\n"
        "3. Change password with the right current password and a valid new password (tick sign out other sessions).\n"
        "4. Open a valid reset link for the same account and set a new password.\n"
        "5. Sign in with the old password.\n"
        "6. Backend: unmark the account, then change the password from /account again."
    ),
    "Suggestion / Expected Behaviour": (
        "Step 2: 'Current password is incorrect' (same as any account). "
        "Steps 3–4: the normal success message is shown, but the stored password does not change and other sessions stay signed in; "
        "the reset link is used up (a second use says the link is invalid or expired). "
        "Step 5: the old password still works. "
        "Step 6: the password changes normally once the account is not core. "
        "Only the core candidate and core employer are flagged; SuperAdmin is not."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P1",
    "Test Status": "Not Run",
    "Comments / Notes": "Owner request 2026-10-08 (shared demo logins). Playwright IS-083 on a test account only.",
    "Phase": "2026-10-08 core account password lock",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate 2; local database",
    "Automation": "Automated",
    "Feature": "Core account password lock",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        for key in ("Issue Summary", "Description", "Suggestion / Expected Behaviour", "Feature"):
            if key in cols:
                ws.cell(row, cols[key]).value = NEW_CASE[key]
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
