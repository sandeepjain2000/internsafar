# -*- coding: utf-8 -*-
"""
Workbook password redaction 2026-10-10.

Older syncs wrote the shared QA sign-in password into Preconditions / Comments / Description cells.
QA passwords live only in .env.local (IP_QA_CORE_PASSWORD for candidates / employers,
IP_QA_SUPERADMIN_PASSWORD for SuperAdmin), so the cells now name the env key instead:

  support@placementhub.online / <password>  ->  support@placementhub.online / <IP_QA_SUPERADMIN_PASSWORD>
  any other occurrence                       ->  <IP_QA_CORE_PASSWORD>

Every sheet and cell is scanned in the canonical workbook and the dated exports in test-cases/.
Throwaway register fixtures that use a longer password (digits after it) are left alone.

  python scripts/sync-internsafar-xlsx-2026-10-10-redact-qa-password.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-10-redact-qa-password.py --apply   # writes the files
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
FILES = sorted((ROOT / "test-cases").glob("InternSafar-Test-Cases*.xlsx"))
OLD = "Admin" + "@123"
SA_RE = re.compile(r"(support@placementhub\.online\s*/\s*)" + re.escape(OLD) + r"(?!\d)")
ANY_RE = re.compile(r"(?<![\w@])" + re.escape(OLD) + r"(?!\d)")


def redact(text: str) -> str:
    text = SA_RE.sub(r"\1<IP_QA_SUPERADMIN_PASSWORD>", text)
    return ANY_RE.sub("<IP_QA_CORE_PASSWORD>", text)


def main(apply: bool) -> None:
    for path in FILES:
        wb = load_workbook(path)
        cells = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and OLD in v:
                        nv = redact(v)
                        if nv != v:
                            cell.value = nv
                            cells += 1
        left = sum(
            1
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
            for v in row
            if isinstance(v, str) and ANY_RE.search(v)
        )
        print(f"{path.name}: {cells} cell(s) redacted, {left} left")
        if apply and cells:
            wb.save(path)
    print("Saved." if apply else "Preview only — re-run with --apply to write.")


if __name__ == "__main__":
    main("--apply" in sys.argv)
