# -*- coding: utf-8 -*-
"""
Add TC-IS-11-021 (one employer-and-candidate offer pipeline, with the four 2026-10-16
screen checks as the done criteria) to InternSafar-Test-Cases.xlsx.

Automated by Playwright IS-088. The employer is whichever of alias1–alias5 already has
an open posting. The candidate is lawsonlclintern+qa-offerflow@gmail.com.
Re-running with --apply updates the row text. It does not reset Test Status.

  python scripts/sync-internsafar-xlsx-2026-10-17-offer-pipeline.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-17-offer-pipeline.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "11 Offers Ratings Completions"
TEMPLATE_ID = "TC-IS-11-020"

NEW_CASE = {
    "ID": "TC-IS-11-021",
    "Module / Section": "11 Offers Ratings Completions",
    "Type": "Positive",
    "Issue Summary": "Offer pipeline: interview edit, letter link, offers list, and search label stay in step",
    "Description": (
        "Employer is the first of alias1, alias3, alias4, alias5, then alias2 that already has a published "
        "posting with no screening questions and an open apply window. Candidate is "
        "lawsonlclintern+qa-offerflow@gmail.com (Neha Kulkarni). Candidate password is <IP_QA_CORE_PASSWORD>. "
        "Alias sign-in is <IP_QA_ALIAS_PASSWORD>. The run does not delete the employer or the posting and does not "
        "change the alias password. It deletes only this candidate's application, offer, and messages.\n"
        "1. Candidate applies. Employer opens that posting and sees Applied, then shortlists.\n"
        "2. Employer opens Schedule interview, saves a time and an https meet link, and the row becomes Interviewing.\n"
        "3. While the row is still Interviewing, the pencil opens Update interview with the saved meet link. "
        "Change the time and the link and save. Do not move the row through Shortlisted to reach that box.\n"
        "4. Send offer shows Upload offer letter and an https link box. Send with an https letter link.\n"
        "5. Candidate Offers, List view, shows Stipend, Starts, Expires, and View letter, Share, Message, Accept, Decline.\n"
        "6. Search Candidates shows Offer sent and no Make offer. After the candidate accepts, the badge is "
        "Offer accepted and Make offer is still absent.\n"
        "7. The employer applicants row is Hired. View letter points at the link that was sent."
    ),
    "Suggestion / Expected Behaviour": (
        "The same application is visible to both people at each step. Status stays Interviewing when the meet "
        "link is edited. The candidate list is not limited to Role, Employer, and Status. Search Candidates "
        "does not keep a Shortlisted badge or a Make offer button after an offer is sent or accepted. "
        "The letter link on the offer is the https link the employer pasted."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P1",
    "Test Status": "Not Run",
    "Comments / Notes": (
        "Written 2026-10-17 from the 2026-10-16 manual offer-acceptance screenshots. Playwright IS-088. "
        "Not run yet. Local or *.vercel.app only; the test skips on internsafar.com. "
        "Out of scope: the colored bulk Shortlist button, endorse, and rate."
    ),
    "Phase": "2026-10-17 offer pipeline",
    "Role(s)": "Employer, Candidate",
    "Preconditions": (
        "One of alias1@placementhub.online, alias3@placementhub.online, alias4@placementhub.online, "
        "alias5@placementhub.online, or alias2@placementhub.online has a published posting with no screening "
        "questions. Candidate lawsonlclintern+qa-offerflow@gmail.com uses <IP_QA_CORE_PASSWORD>. "
        "The alias signs in with <IP_QA_ALIAS_PASSWORD>."
    ),
    "Automation": "Automated",
    "Feature": "Offer pipeline",
}

KEEP_ON_UPDATE = {"Test Status", "Date Verified", "Verified By", "ID"}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        for key, value in NEW_CASE.items():
            if key in KEEP_ON_UPDATE or key not in cols:
                continue
            ws.cell(row, cols[key]).value = value
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return

    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
