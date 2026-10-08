# -*- coding: utf-8 -*-
"""
Add TC-IS-08-005 (My Applications: sort, status tabs, metric counts, search, column filters and which
rows can be withdrawn) to InternSafar-Test-Cases.xlsx.

Automated by Playwright IS-084 on a fixed set of seven applications built with the real
src/lib/ipApplicationPresentation.js (one per status, plus an Awaiting Review row on a closed posting).
TC-IS-08-001 / TC-IS-08-002 are covered on real data by IS-085 and the checklist runner (CAND-AP-1/2).
Re-running with --apply updates the row if it already exists.

  python scripts/sync-internsafar-xlsx-2026-10-08-applications-list.py           # preview only
  python scripts/sync-internsafar-xlsx-2026-10-08-applications-list.py --apply   # writes workbook
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "test-cases" / "InternSafar-Test-Cases.xlsx"
SHEET = "08 Candidate Pipeline UX"
TEMPLATE_ID = "TC-IS-08-004"

NEW_CASE = {
    "ID": "TC-IS-08-005",
    "Module / Section": "08 Candidate Pipeline UX",
    "Type": "Functional",
    "Issue Summary": "My Applications sort, status tabs, metric counts, search, column filters and withdraw availability",
    "Description": (
        "1. Sign in as the test candidate with applications in several statuses and open /candidate/applications.\n"
        "2. Read the four cards: Total Submitted, In Review, Interviews Scheduled, Offers Received.\n"
        "3. Sort by: Latest First, Oldest First, Highest match, Status.\n"
        "4. Click each status tab (Awaiting Review … Withdrawn), then All Applications.\n"
        "5. Search by company, then by role in capitals, then by text that matches nothing → Clear Status Filters.\n"
        "6. Filters → Employer; add Status; Clear. Next step = Attend the scheduled interview; Clear. Applied from a date; Clear.\n"
        "7. Check the Withdraw button on each row; open View details on an interview row and on an Awaiting Review row."
    ),
    "Suggestion / Expected Behaviour": (
        "Step 2: Total = all applications; In Review = Under Review + Awaiting Review on postings still open "
        "(not closed/expired); Interviews = Interview Scheduled; Offers = Offer Received + Hired. "
        "Step 3: newest first / oldest first / highest match first / alphabetical by status label. "
        "Step 4: each tab shows only its status (Offer Received tab includes Hired/Completed; Rejected includes Offer Declined); "
        "the active tab is highlighted; All Applications shows the total in brackets; cards do not change. "
        "Step 5: search matches role or company, any case; no match shows 'No applications found'; Clear restores everything. "
        "Step 6: each filter narrows the list, the Filters chip counts active filters, Clear restores the list. "
        "Step 7: Withdraw is enabled only for Awaiting Review rows; the interview row's details show 'Attend interview <date, time>' "
        "and no Withdraw button; the Awaiting Review row's details show Withdraw application."
    ),
    "Doc Status": "Ready for QA",
    "Severity / Priority": "P2",
    "Test Status": "Not Run",
    "Comments / Notes": "Added 2026-10-08 (My Applications automation gap). Playwright IS-084 (fixed data, real status logic).",
    "Phase": "2026-10-08 applications list automation",
    "Role(s)": "Candidate",
    "Preconditions": "test candidate",
    "Automation": "Automated",
    "Feature": "My Applications list",
}


def find_header(ws):
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r, {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    return None, {}


def main(apply: bool) -> None:
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    hr, cols = find_header(ws)
    if not hr:
        raise SystemExit(f"No header row in {SHEET}")

    rows = {str(ws.cell(r, cols["ID"]).value or "").strip(): r for r in range(hr + 1, ws.max_row + 1)}
    if NEW_CASE["ID"] in rows:
        row = rows[NEW_CASE["ID"]]
        for key in ("Issue Summary", "Description", "Suggestion / Expected Behaviour", "Feature"):
            if key in cols:
                ws.cell(row, cols[key]).value = NEW_CASE[key]
        print(f"{'Updating' if apply else 'Would update'} {NEW_CASE['ID']} text at row {row} of '{SHEET}'.")
        if apply:
            wb.save(XLSX)
            print(f"Saved {XLSX.name}")
        return
    template = rows.get(TEMPLATE_ID)
    if not template:
        raise SystemExit(f"Template row {TEMPLATE_ID} not found")

    target = max(rows.values()) + 1
    for key, c in cols.items():
        cell = ws.cell(target, c)
        cell._style = copy(ws.cell(template, c)._style)
        cell.value = NEW_CASE.get(key)
    print(f"{'Adding' if apply else 'Would add'} {NEW_CASE['ID']} at row {target} of '{SHEET}'.")
    if apply:
        wb.save(XLSX)
        print(f"Saved {XLSX.name}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
